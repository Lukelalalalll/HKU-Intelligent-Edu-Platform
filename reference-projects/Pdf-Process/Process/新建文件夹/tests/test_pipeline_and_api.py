import hashlib
import json
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
import pytest

from app.main import app
from app.db import SessionLocal
from app.models import Document, ProcessingJob
from app.services.pipeline import process_document
from app.services.retrieval import retrieve
from fastapi.testclient import TestClient

FIXTURE_PDF = Path(__file__).parents[2] / 'data' / 'documents' / '1' / 'source.pdf'

def make_doc(tmp_path, source, file_type):
    staging = tmp_path / 'staging'; staging.mkdir(parents=True, exist_ok=True)
    target = staging / f'source.{file_type}'; shutil.copy2(source, target)
    if file_type == 'pdf': target.write_bytes(target.read_bytes() + b'\n' + str(tmp_path).encode())
    db = SessionLocal()
    data = target.read_bytes(); doc = Document(filename=target.name, file_type=file_type, sha256=hashlib.sha256(data).hexdigest()); db.add(doc); db.commit(); db.refresh(doc)
    job = ProcessingJob(document_id=doc.id); db.add(job); db.commit(); db.refresh(job)
    root = tmp_path / 'documents' / str(doc.id); root.mkdir(parents=True, exist_ok=True); shutil.copy2(target, root / target.name)
    return db, doc, job, root

def test_pdf_structured_fallback(tmp_path, monkeypatch):
    monkeypatch.setattr('app.config.settings.data_dir', tmp_path)
    db, doc, job, root = make_doc(tmp_path, FIXTURE_PDF, 'pdf')
    process_document(db, doc, job)
    manifest = json.loads((root / 'manifest.json').read_text())
    assert doc.status == 'ready' and manifest['successful_pages'] == 5
    assert '<!-- page: 1 -->' in (root / 'document.md').read_text(encoding='utf-8')
    blocks = json.loads((root / 'pages/page-001.json').read_text(encoding='utf-8'))['blocks']
    assert all('bbox_raw' in b and 'confidence' in b and 'page_number' in b for b in blocks)
    assert all('\\=' not in b.get('markdown','') for b in blocks)

def test_pptx_title_code_and_slide_chunks(tmp_path, monkeypatch):
    monkeypatch.setattr('app.config.settings.data_dir', tmp_path)
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = 'Demo'; slide.placeholders[1].text = 'import pandas as pd\nscore = df.sum()'
    source = tmp_path / 'sample.pptx'; prs.save(source)
    db, doc, job, root = make_doc(tmp_path, source, 'pptx'); process_document(db, doc, job)
    slide_json = json.loads((root / 'slides/slide-001.json').read_text()); assert slide_json['title'] == 'Demo'; assert slide_json['code_blocks']
    chunks = [json.loads(x) for x in (root / 'chunks.jsonl').read_text().splitlines()]; assert chunks and chunks[0]['document_id'] == doc.id and chunks[0]['slide_number'] == 1
    assert json.loads((root / 'manifest.json').read_text())['rendering']['status'] in {'unavailable','ready'}

def test_upload_and_chat_mock():
    client = TestClient(app)
    assert client.get('/health').status_code == 200
    assert client.post('/api/v1/documents', files={'files': ('bad.txt', b'x', 'text/plain')}).status_code == 400
    conversation = client.post('/api/v1/chat/conversations', json={'document_ids': [1]}).json()
    response = client.post(f"/api/v1/chat/conversations/{conversation['id']}/messages", json={'content': '公式', 'provider': 'mock', 'document_ids': [1]})
    assert response.status_code == 200 and response.json()['test_mode'] is True
    stream = client.post(f"/api/v1/chat/conversations/{conversation['id']}/stream", json={'content': '代码', 'provider': 'mock', 'document_ids': [1]})
    assert stream.status_code == 200 and 'text/event-stream' in stream.headers['content-type'] and 'token' in stream.text

def test_missing_key_returns_503():
    client = TestClient(app); conversation = client.post('/api/v1/chat/conversations', json={}).json()
    response = client.post(f"/api/v1/chat/conversations/{conversation['id']}/messages", json={'content': 'hello', 'provider': 'openai-compatible'})
    assert response.status_code == 503 and 'API_KEY' in response.text

def test_provider_validation_and_retrieval_shape():
    from app.services.ai_provider import provider_config, ProviderError
    with pytest.raises(ProviderError): provider_config(provider='unknown', api_key='x')
    client = TestClient(app); result = client.get('/api/v1/search', params={'q': 'Python', 'top_k': 1}).json()
    assert result and {'page', 'confidence', 'filename', 'chunk_id'} <= set(result[0])
