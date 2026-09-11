from __future__ import annotations
import hashlib, json, re, shutil, subprocess, tempfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
import fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy.orm import Session
from ..models import Document, ProcessingJob, Page, Chunk
from ..config import settings
from .mineru_runner import run_mineru, find_middle_json
from .chunking import build_chunks

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1024 * 1024), b''): h.update(b)
    return h.hexdigest()

def _safe_error(value):
    value = str(value or '')
    if settings.ai_api_key:
        value = value.replace(settings.ai_api_key, '[REDACTED]')
    return value[-4000:]

def process_document(db: Session, doc: Document, job: ProcessingJob):
    root = settings.documents_dir / str(doc.id); root.mkdir(parents=True, exist_ok=True)
    ext = (doc.file_type or Path(doc.filename).suffix.lstrip('.')).lower()
    src = root / f'source.{ext}'
    job.status = doc.status = 'processing'; job.stage = 'extracting'; db.commit()
    try:
        if ext == 'pptx':
            process_pptx(db, doc, job, src, root); return
        process_pdf(db, doc, job, src, root)
    except Exception as exc:
        job.status = doc.status = 'failed'; job.error = doc.error = _safe_error(exc); db.commit(); raise

def _write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')

def process_pdf(db, doc, job, src, root):
    pdf = fitz.open(src); total = pdf.page_count; doc.pages = total; job.total_pages = total; db.commit(); pdf.close()
    mineru_error = None; pages = []
    try:
        out = root / 'mineru'; run_mineru(src, out); middle = find_middle_json(out)
        if middle and middle.exists():
            raw = json.loads(middle.read_text(encoding='utf-8'))
            for i, page in enumerate(raw.get('pdf_info', raw.get('pages', [])), 1): pages.append(_miner_page(page, i))
    except Exception as exc:
        mineru_error = _safe_error(exc)
    if not pages: pages = _extract_native_pages(src, root)
    pages = _merge_code_blocks(pages)
    page_dir = root / 'pages'; (root / 'assets').mkdir(exist_ok=True); (root / 'logs').mkdir(exist_ok=True)
    successful = failed = low_conf = 0; db.query(Page).filter_by(document_id=doc.id).delete(); db.query(Chunk).filter_by(document_id=doc.id).delete(); db.commit()
    all_chunks = []
    for p in pages:
        try:
            _write_json(page_dir / f"page-{p['page']:03d}.json", p)
            db.add(Page(document_id=doc.id, page_number=p['page'], width=p['width'], height=p['height'], payload=p))
            successful += 1; low_conf += sum(1 for b in p['blocks'] if b.get('confidence', 1) < .7)
            all_chunks.extend(build_chunks([p], document_id=doc.id, source_file=doc.filename))
        except Exception as exc:
            failed += 1; job.log = (job.log or '') + f"\nPage {p.get('page')} failed: {_safe_error(exc)}"
        job.processed_pages = successful + failed; job.progress = job.processed_pages / max(total, 1) * 100; job.status = doc.status = 'partial_ready'; db.commit()
    for c in all_chunks:
        db.add(Chunk(document_id=doc.id, page=c['page'], section=c.get('section'), chunk_type=c['chunk_type'], content=c['content'], metadata_json=c['metadata']))
    db.commit()
    md = _pages_to_markdown(pages); (root / 'document.md').write_text(md, encoding='utf-8')
    (root / 'chunks.jsonl').write_text(''.join(json.dumps(c, ensure_ascii=False) + '\n' for c in all_chunks), encoding='utf-8')
    manifest = {
        'document_id': doc.id, 'filename': doc.filename, 'file_type': 'pdf', 'sha256': doc.sha256,
        'pages': total, 'successful_pages': successful, 'failed_pages': failed,
        'low_confidence_pages': sum(1 for p in pages if any(b.get('confidence', 1) < .7 for b in p['blocks'])),
        'source': 'mineru' if not mineru_error and any(b.get('source') == 'mineru' for p in pages for b in p['blocks']) else 'native_pdf_fallback',
        'mineru_parse_method': settings.mineru_parse_method, 'mineru_error': mineru_error,
        'omissions': [f'page {p["page"]}' for p in pages if not p.get('blocks')],
        'stats': _page_stats(pages), 'artifacts': ['source.pdf', 'document.md', 'manifest.json', 'chunks.jsonl'],
        'page_files': [f'pages/page-{p["page"]:03d}.json' for p in pages],
    }
    _write_json(root / 'manifest.json', manifest)
    job.status = doc.status = 'ready' if failed == 0 else 'partial_ready'; job.stage = 'ready'; job.progress = 100; job.finished_at = datetime.now(timezone.utc); db.commit()

def _miner_page(page, number):
    blocks = []
    for n, b in enumerate(page.get('para_blocks', page.get('blocks', [])), 1):
        raw = b.get('bbox') or [0, 0, 1, 1]; text = b.get('text') or b.get('content') or ''; typ = _normalize_block_type(b.get('type'), text)
        latex = b.get('latex') or b.get('formula') or None; confidence = float(b.get('confidence', .92))
        blocks.append(_block(f'p{number}-b{n}', typ, text, raw, page.get('page_width', 1), page.get('page_height', 1), 'mineru', confidence, latex))
    return {'page': number, 'width': page.get('page_width', 0), 'height': page.get('page_height', 0), 'blocks': blocks}

def _extract_native_pages(src, root):
    pdf = fitz.open(src); pages = []
    for i, page in enumerate(pdf, 1):
        blocks = []
        for n, raw in enumerate(page.get_text('dict').get('blocks', []), 1):
            if raw.get('type') != 0: continue
            lines = []; fonts = []; sizes = []; gaps = []
            for line in raw.get('lines', []):
                fonts += [s.get('font', '') for s in line.get('spans', [])]; sizes += [s.get('size', 0) for s in line.get('spans', [])]
                value = ''.join(s.get('text', '') for s in line.get('spans', [])).strip()
                if value: lines.append(value)
                gaps.append(line.get('bbox', [0, 0, 0, 0])[3])
            text = '\n'.join(lines).strip()
            if not text: continue
            typ = _normalize_block_type('', text)
            if any(x in f.lower() for f in fonts for x in ('mono', 'courier', 'cmtt')): typ = 'code'
            if _formula_candidate(text): typ = 'formula'
            if _looks_like_heading(text, sizes, page.rect.height): typ = 'heading'
            confidence = .82 if typ == 'formula' else .96
            b = _block(f'p{i}-b{n}', typ, text, raw.get('bbox', [0, 0, page.rect.width, page.rect.height]), page.rect.width, page.rect.height, 'native_pdf', confidence, text if typ == 'formula' else None, fonts, max(sizes or [0]))
            b['line_spacing'] = max(gaps) - min(gaps) if len(gaps) > 1 else 0
            if typ == 'formula' and confidence < .9:
                b['needs_visual_verification'] = True
                try:
                    clip = fitz.Rect(raw.get('bbox', [0, 0, page.rect.width, page.rect.height])); image = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), clip=clip)
                    shot = root / 'assets' / f'page-{i:03d}-formula-{n:03d}.png'; shot.parent.mkdir(exist_ok=True); image.save(str(shot)); b['image'] = str(shot.relative_to(root))
                except Exception: pass
            blocks.append(b)
        # Image regions are evidence even when OCR/text extraction returns no text.
        try:
            for n, info in enumerate(page.get_image_info(full=True), 1):
                rect = info.get('bbox') or [0, 0, page.rect.width, page.rect.height]
                blocks.append(_block(f'p{i}-image-{n}', 'image', '', rect, page.rect.width, page.rect.height, 'native_pdf_image', .9))
        except Exception:
            pass
        pages.append({'page': i, 'width': page.rect.width, 'height': page.rect.height, 'blocks': blocks})
    pdf.close(); return pages

def _block(block_id, typ, text, raw, width, height, source, confidence, latex=None, fonts=None, size=0):
    x0, y0, x1, y1 = [float(v) for v in raw]
    clean = _clean_text(text) if typ not in ('code', 'formula') else text
    return {'id': block_id, 'page_number': int(block_id.split('-')[0][1:]), 'type': typ, 'raw_text': text, 'text': clean, 'markdown': _block_markdown(typ, clean, latex), 'latex': latex if typ == 'formula' else None, 'inline_formula_candidates': re.findall(r'(?<!\\$)\\$([^$]+)\\$', text or ''), 'bbox_raw': [x0, y0, x1, y1], 'bbox': [max(0, min(1, x0 / max(width, 1))), max(0, min(1, y0 / max(height, 1))), max(0, min(1, x1 / max(width, 1))), max(0, min(1, y1 / max(height, 1)))], 'font': fonts or [], 'font_size': size, 'confidence': confidence, 'source': source, 'needs_visual_verification': False}

def _clean_text(value):
    value = str(value or '')
    value = re.sub(r'-\s*\n\s*', '', value)
    return re.sub(r'(?<!\n)[ \t]*\n[ \t]*', ' ', value).strip()

def _sanitize_latex(value):
    value = str(value or '').strip().replace('\\=', '=')
    return value.replace('$$', '').strip()

def _normalize_block_type(raw, text):
    value = str(raw or '').lower(); text = text or ''
    if any(x in value for x in ('equation', 'formula', 'interline')) or _formula_candidate(text): return 'formula'
    if 'table' in value: return 'table'
    if 'image' in value or 'figure' in value: return 'image'
    if 'heading' in value or 'title' in value: return 'heading'
    if 'code' in value or _looks_like_code(text): return 'code'
    if re.match(r'^\s*(?:[-*•]|\d+[.)])\s+', text): return 'list'
    return 'text'

def _looks_like_code(text):
    lines = text.replace('\\n', '\n').splitlines(); joined = '\n'.join(lines)
    if len(lines) < 2: return False
    syntax = sum(bool(re.search(r'^\s*(?:def |class |from |import |return |for |while |if |SELECT |INSERT |const |public |System\\.out)', x, re.I)) for x in lines)
    operators = sum(bool(re.search(r'==|:=|=>|;|\{|\}|\(.*\):', x)) for x in lines)
    return syntax >= 2 or (syntax >= 1 and (len(lines) >= 2 or operators >= 1)) or (operators >= 3 and any('    ' in x or '\t' in x for x in lines))

def _formula_candidate(text):
    return len(text) < 240 and (bool(re.search(r'[∑∫√∞≤≥≈≠∈×÷]', text)) or bool(re.search(r'\b(?:softmax|Similarity|argmax|log|exp)\s*\(', text)))

def _looks_like_heading(text, sizes, page_height):
    return len(text) < 160 and (len(sizes) > 0 and max(sizes) >= 14 or re.match(r'^(?:chapter|section|question|assignment)\b', text, re.I))

def _merge_code_blocks(pages):
    for page in pages:
        merged = []
        for block in page.get('blocks', []):
            if merged and block.get('type') == merged[-1].get('type') == 'code':
                merged[-1]['text'] += '\n' + block['text']; merged[-1]['raw_text'] = merged[-1]['text']; merged[-1]['markdown'] = _block_markdown('code', merged[-1]['text'])
            else: merged.append(block)
        page['blocks'] = merged
    return pages

def _block_markdown(typ, text, latex=None):
    value = (latex or text or '').strip()
    if typ == 'formula':
        value = _sanitize_latex(value)
        return f'$$\n{value}\n$$' if value else ''
    if typ == 'code': return '```text\n' + value + '\n```'
    if typ == 'heading': return '### ' + value
    if typ == 'list': return value
    return value

def _pages_to_markdown(pages):
    return '\n\n'.join(f'<!-- page: {p["page"]} -->\n\n' + '\n\n'.join(b.get('markdown') or b.get('text', '') for b in p.get('blocks', [])) for p in pages)

def _page_stats(pages):
    blocks = [b for p in pages for b in p.get('blocks', [])]
    return {'block_count': len(blocks), 'formula_count': sum(b.get('type') == 'formula' for b in blocks), 'code_count': sum(b.get('type') == 'code' for b in blocks), 'table_count': sum(b.get('type') == 'table' for b in blocks), 'image_count': sum(b.get('type') == 'image' for b in blocks)}

def process_pptx(db, doc, job, src, root):
    prs = Presentation(str(src)); total = len(prs.slides); job.total_pages = total; doc.pages = total; db.commit()
    slide_dir = root / 'slides'; page_dir = root / 'pages'; slide_dir.mkdir(exist_ok=True); page_dir.mkdir(exist_ok=True); (root / 'logs').mkdir(exist_ok=True)
    pages = []; chunks = []; failed = 0
    for number, slide in enumerate(prs.slides, 1):
        try:
            page = _parse_slide(prs, slide, number, doc.filename)
            pages.append(page); _write_json(slide_dir / f'slide-{number:03d}.json', page); _write_json(page_dir / f'page-{number:03d}.json', page)
            chunks += build_chunks([page], document_id=doc.id, source_file=doc.filename)
        except Exception as exc:
            failed += 1; job.log = (job.log or '') + f"\nSlide {number} failed: {_safe_error(exc)}"
        job.processed_pages = number; job.progress = number / max(total, 1) * 100; job.status = doc.status = 'partial_ready'; db.commit()
    db.query(Page).filter_by(document_id=doc.id).delete(); db.query(Chunk).filter_by(document_id=doc.id).delete()
    for p in pages: db.add(Page(document_id=doc.id, page_number=p['page'], width=p['width'], height=p['height'], payload=p))
    for c in chunks: db.add(Chunk(document_id=doc.id, page=c['page'], section=c.get('section'), chunk_type=c['chunk_type'], content=c['content'], metadata_json=c['metadata']))
    db.commit(); (root / 'document.md').write_text(_pages_to_markdown(pages), encoding='utf-8'); (root / 'chunks.jsonl').write_text(''.join(json.dumps(c, ensure_ascii=False) + '\n' for c in chunks), encoding='utf-8')
    render = _render_pptx(src, root, total)
    manifest = {'document_id': doc.id, 'filename': doc.filename, 'file_type': 'pptx', 'sha256': doc.sha256, 'pages': total, 'successful_pages': len(pages), 'failed_pages': failed, 'low_confidence_pages': sum(1 for p in pages if p.get('confidence', 1) < .7), 'source': 'python-pptx', 'rendering': render, 'artifacts': ['source.pptx', 'document.md', 'manifest.json', 'chunks.jsonl'], 'slide_files': [f'slides/slide-{p["page"]:03d}.json' for p in pages], 'page_files': [f'pages/page-{p["page"]:03d}.json' for p in pages]}
    _write_json(root / 'manifest.json', manifest); job.status = doc.status = 'ready' if failed == 0 else 'partial_ready'; job.stage = 'ready'; job.progress = 100; job.finished_at = datetime.now(timezone.utc); db.commit()

def _parse_slide(prs, slide, number, filename):
    blocks = []; tables = []; images = []; charts = []; formulas = []; warnings = []
    for index, shape in enumerate(slide.shapes, 1):
        bbox = [shape.left / prs.slide_width, shape.top / prs.slide_height, (shape.left + shape.width) / prs.slide_width, (shape.top + shape.height) / prs.slide_height]
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]; tables.append({'bbox': bbox, 'rows': rows}); blocks.append({'id': f'p{number}-b{index}', 'type': 'table', 'text': '\n'.join(' | '.join(r) for r in rows), 'markdown': '\n'.join(' | '.join(r) for r in rows), 'bbox': bbox, 'bbox_raw': [shape.left, shape.top, shape.left + shape.width, shape.top + shape.height], 'source': 'pptx', 'confidence': .96}); continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE: images.append({'bbox': bbox, 'name': shape.name}); continue
        if shape.shape_type == MSO_SHAPE_TYPE.CHART: charts.append({'bbox': bbox, 'description': getattr(shape.chart, 'chart_type', 'chart')}); continue
        if not getattr(shape, 'has_text_frame', False): continue
        text = '\n'.join(p.text for p in shape.text_frame.paragraphs).strip().replace('\\n', '\n')
        if not text: continue
        is_title = getattr(shape, 'is_placeholder', False) and getattr(shape.placeholder_format, 'type', None) in (1, 2)
        typ = 'heading' if is_title else ('formula' if _formula_candidate(text) else ('code' if _looks_like_code(text) else 'text')); block = {'id': f'p{number}-b{index}', 'type': typ, 'raw_text': text, 'text': text, 'markdown': _block_markdown(typ, text), 'latex': text if typ == 'formula' else None, 'bbox': bbox, 'bbox_raw': [shape.left, shape.top, shape.left + shape.width, shape.top + shape.height], 'font': [], 'font_size': 0, 'confidence': .86 if typ == 'formula' else .95, 'source': 'pptx', 'needs_visual_verification': typ == 'formula'}; blocks.append(block)
    notes = ''
    try:
        notes_slide = slide.notes_slide
        if hasattr(notes_slide, 'notes_text_frame'):
            notes = '\n'.join(s.text for s in notes_slide.notes_text_frame.paragraphs).strip()
        if not notes:
            notes = '\n'.join(getattr(shape, 'text', '') for shape in notes_slide.shapes if getattr(shape, 'has_text_frame', False)).strip()
    except Exception: pass
    title = next((b['text'] for b in blocks if b['type'] == 'heading'), '')
    formulas = [b['latex'] for b in blocks if b.get('type') == 'formula' and b.get('latex')]
    return {'page': number, 'slide_number': number, 'width': prs.slide_width, 'height': prs.slide_height, 'title': title, 'body_text': '\n\n'.join(b['text'] for b in blocks if b['type'] != 'heading'), 'notes': notes, 'tables': tables, 'images': images, 'charts': charts, 'formulas': formulas, 'code_blocks': [b['text'] for b in blocks if b['type'] == 'code'], 'blocks': blocks, 'confidence': min([b.get('confidence', 1) for b in blocks] or [1]), 'warnings': warnings}

def _render_pptx(src, root, total):
    command = settings.render_command or shutil.which('soffice') or shutil.which('libreoffice')
    if not command: return {'status': 'unavailable', 'reason': 'LibreOffice/soffice not found; retained structured JSON only', 'rendered_pages': 0}
    try:
        out = root / 'rendered'; out.mkdir(exist_ok=True); subprocess.run([command, '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(src)], capture_output=True, text=True, timeout=300, check=True)
        pdf = next(out.glob('*.pdf')); image_dir = root / 'slides'; subprocess.run(['pdftoppm', '-png', '-r', str(settings.render_dpi), str(pdf), str(image_dir / 'slide')], capture_output=True, text=True, timeout=300, check=True)
        for i, path in enumerate(sorted(image_dir.glob('slide-*.png')), 1): path.rename(image_dir / f'slide-{i:03d}.png')
        return {'status': 'ready', 'rendered_pages': total, 'tool': command}
    except Exception as exc: return {'status': 'failed', 'reason': _safe_error(exc), 'rendered_pages': 0}
