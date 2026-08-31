from types import SimpleNamespace

from pptx import Presentation

from app.services.ppt_export import export_project_pptx


def test_export_project_creates_native_powerpoint_objects(tmp_path):
    page = SimpleNamespace(
        id="page-1",
        sort_order=0,
        title="导出测试",
        section_title="测试章节",
        bullets_json=["可导出的要点"],
        design_document_json={
            "version": 2,
            "theme": {"background": "#f2c94c"},
            "elements": [{"id": "title", "type": "title", "x": 7, "y": 8, "w": 86, "h": 12, "text": "导出测试"}],
        },
        speaker_notes="",
    )
    project = SimpleNamespace(theme_id="light-academic", layout_assignments={}, pages=[page])
    output = tmp_path / "courseware.pptx"

    export_project_pptx(project, output)

    assert output.exists() and output.stat().st_size > 0
    presentation = Presentation(str(output))
    assert len(presentation.slides) == 1
    assert presentation.slides[0].shapes[0].text == "导出测试"


def test_export_preserves_decorative_shapes_and_alias_fields(tmp_path):
    page = SimpleNamespace(
        id="page-visuals",
        sort_order=0,
        title="视觉元素",
        section_title="",
        bullets_json=[],
        design_document_json={
            "theme": {"background": "#eeeeee", "surface": "#ffffff", "accent": "#2f80ed"},
            "elements": [
                {"type": "panel", "x": 10, "y": 10, "w": 80, "h": 70, "background_color": "#ffffff", "radius": 4},
                {"type": "divider", "x": 15, "y": 40, "w": 40, "h": 0.5, "stroke_color": "#2f80ed", "radius": 1},
                {"type": "circle", "x": 80, "y": 20, "w": 8, "h": 8, "fill_color": "#27ae60"},
                {"type": "body", "x": 15, "y": 20, "w": 50, "h": 8, "text": "保留文字"},
            ],
        },
        speaker_notes="",
    )
    project = SimpleNamespace(theme_id="light-academic", layout_assignments={}, pages=[page])
    output = tmp_path / "visuals.pptx"

    export_project_pptx(project, output)

    slide = Presentation(str(output)).slides[0]
    # Three native visual objects plus the text box; decorative elements must
    # not disappear merely because they do not contain text.
    assert len(slide.shapes) == 4
    assert sum(1 for shape in slide.shapes if shape.has_text_frame and shape.text == "") == 3
