from __future__ import annotations

import json
import hashlib
import re
import shutil
import threading
import base64
import mimetypes
import ipaddress
import socket
from urllib.parse import urlparse
from datetime import datetime, timezone
from concurrent.futures import ThreadPoolExecutor
from functools import lru_cache
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.services.ai_gateway import AIGateway
from app.models import (
    PptAgentEvent, PptAgentRun, PptExportJob, PptOutlineVersion, PptPage,
    PptProject, PptRequirement, PptSourceChunk,
    PptSourceCollection, PptSourceDocument, User, PptMessage, PptCheckpoint, PptDocumentVersion, PptGenerationJob,
)
from app.db.session import SessionLocal
from app.services.ppt_theme import preview_filename, get_theme, list_layouts, validate_theme_pack
from app.services.browser_image_search import BrowserImageSearch, build_query
from app.services.image_cache import download_image, validate_public_url
from app.services.visual_search import extract_keywords, query_text, rank_candidates, stable_asset_id
from app.services.clip_ranker import rerank


STAGES = ("init", "outline", "visual", "theme", "layout", "design", "export")
STATUSES = ("empty", "ready", "running", "confirmed", "stale", "failed")
generation_executor = ThreadPoolExecutor(max_workers=max(1, min(int(getattr(settings, "ppt_generation_executor_workers", 4)), 8)))
visual_executor = ThreadPoolExecutor(max_workers=max(2, min(int(getattr(settings, "ppt_visual_executor_workers", 6)), 12)))
_SOURCE_CONTEXT_CACHE: dict[str, list[dict[str, Any]]] = {}
_SOURCE_CONTEXT_LOCK = threading.Lock()
_GENERATION_LOCKS: dict[str, threading.Lock] = {}
_GENERATION_LOCKS_GUARD = threading.Lock()


def _mask(value: str) -> str:
    if not value:
        return ""
    return f"{value[:4]}••••{value[-4:]}" if len(value) > 8 else "••••••••"


def _cipher():
    try:
        from cryptography.fernet import Fernet
        import base64, hashlib
        key = base64.urlsafe_b64encode(hashlib.sha256(settings.jwt_secret_key.encode()).digest())
        return Fernet(key)
    except ImportError:
        # Keep the app runnable before optional wheels are installed; production
        # requirements include cryptography/Fernet.
        import base64, hashlib
        key = hashlib.sha256(settings.jwt_secret_key.encode()).digest()
        class LocalCipher:
            def encrypt(self, value: bytes) -> bytes:
                return base64.urlsafe_b64encode(bytes(v ^ key[i % len(key)] for i, v in enumerate(value)))
            def decrypt(self, value: bytes) -> bytes:
                raw = base64.urlsafe_b64decode(value)
                return bytes(v ^ key[i % len(key)] for i, v in enumerate(raw))
        return LocalCipher()


def encrypt_api_key(value: str) -> str:
    return _cipher().encrypt(value.encode()).decode()


def decrypt_api_key(value: str) -> str:
    if not value:
        return ""
    return _cipher().decrypt(value.encode()).decode()


class ProviderGateway(AIGateway):
    """Backward-compatible PPT gateway backed by the platform AI mapping."""
    def __init__(self, db: Session, user_id: str | None = None, business_code: str = "teacher_ppt_agent"):
        super().__init__(db, business_code)
        from types import SimpleNamespace
        self.config = SimpleNamespace(model=self.model, embedding_model=(self.binding.embedding_model if self.binding else ""))
@lru_cache(maxsize=32)
def _cached_openai_client(api_key: str, base_url: str, timeout: int):
    from openai import OpenAI
    return OpenAI(api_key=api_key, base_url=base_url, timeout=timeout)


class PptAgentService:
    def __init__(self, db: Session, user: User):
        self.db, self.user = db, user

    def project(self, project_id: str) -> PptProject:
        item = self.db.scalar(select(PptProject).where(PptProject.id == project_id, PptProject.owner_id == self.user.id))
        if not item:
            raise HTTPException(404, "PPT 项目不存在")
        return item

    def page(self, project_id: str, page_id: str) -> PptPage:
        item = self.db.scalar(select(PptPage).where(PptPage.id == page_id, PptPage.project_id == project_id))
        if not item or item.project.owner_id != self.user.id:
            raise HTTPException(404, "页面不存在")
        return item

    def serialize_project(self, p: PptProject) -> dict[str, Any]:
        pages = sorted(p.pages, key=lambda item: item.sort_order)
        if p.theme_id:
            base_theme = get_theme(p.theme_id)
            theme = validate_theme_pack({**base_theme, **(p.theme_config or {}), "colors": {**(base_theme.get("colors") or {}), **((p.theme_config or {}).get("colors") or {})}, "tokens": {**(base_theme.get("tokens") or {}), **((p.theme_config or {}).get("tokens") or {})}})
        else:
            theme = p.theme_config or {}
        return {"id": p.id, "title": p.title, "request_text": p.request_text, "course_id": p.course_id, "current_stage": p.current_stage, "status": p.status, "active_page_id": p.active_page_id, "latest_checkpoint_code": p.latest_checkpoint_code, "page_count_target": p.page_count_target, "theme_id": p.theme_id, "theme_config": theme, "layout_assignments": p.layout_assignments or {}, "design_status": p.design_status, "page_count": len(pages), "cover_preview_url": preview_filename(p.id, pages[0].id) if pages else None, "created_at": p.created_at.isoformat(), "updated_at": p.updated_at.isoformat()}

    def delete_project(self, project_id: str) -> None:
        project = self.project(project_id)
        project_root = settings.ppt_storage_path / project.id
        export_path = settings.ppt_storage_path / f"{project.id}.pptx"
        self.db.delete(project)
        self.db.commit()
        if project_root.is_dir():
            shutil.rmtree(project_root)
        if export_path.is_file():
            export_path.unlink()

    def serialize_page(self, p: PptPage) -> dict[str, Any]:
        statuses = dict(p.statuses_json or {})
        document = p.design_document_json or p.draft_document_json
        visual_plan = p.visual_plan_json or {}
        return {"id": p.id, "project_id": p.project_id, "section_title": p.section_title, "sort_order": p.sort_order, "title": p.title, "bullets": p.bullets_json or [], "statuses": statuses, "search_queries": p.search_queries_json or [], "summary_md": p.summary_md, "citations": p.citations_json or [], "document": document, "document_revision": p.document_revision, "current_document_version_id": p.current_document_version_id, "speaker_notes": p.speaker_notes, "page_role": p.page_role, "content_plan": p.content_plan_json or {}, "visual_plan": visual_plan, "asset_manifest": visual_plan.get("asset_manifest", []), "image_slots": visual_plan.get("image_slots", []), "selected_asset_ids": visual_plan.get("selected_asset_ids", []), "preview_url": preview_filename(p.project_id, p.id) if document else None, "layout_id": (p.project.layout_assignments or {}).get(p.id), "created_at": p.created_at.isoformat(), "updated_at": p.updated_at.isoformat()}

    def list_projects(self) -> list[dict[str, Any]]:
        return [self.serialize_project(p) for p in self.db.scalars(select(PptProject).where(PptProject.owner_id == self.user.id).order_by(PptProject.updated_at.desc()))]

    def create_project(self, title: str, request_text: str, course_id: str | None = None) -> dict[str, Any]:
        p = PptProject(owner_id=self.user.id, title=title.strip() or "未命名课件", request_text=request_text.strip(), course_id=course_id)
        self.db.add(p); self.db.flush()
        self.db.add(PptRequirement(project_id=p.id))
        self._message(p, "user", p.request_text, payload={"intent_type": "initial_request"})
        self._message(p, "assistant", "好的，我先帮你把课件需求梳理清楚。先确认一下：这套 PPT 主要面向什么对象？例如本科生、研究生、教师培训或企业客户。", payload={
            "intent_type": "clarify_requirements",
            "question": {"code": "audience", "label": "这套 PPT 主要面向什么对象？", "options": ["本科生", "研究生", "教师培训", "企业客户"]},
            "ready_to_outline": False,
        })
        self._event(p, "project.created", {"stage": "init"})
        self.db.commit(); self.db.refresh(p)
        return self.serialize_project(p)

    def list_messages(self, project_id: str, page_id: str | None = None) -> list[dict[str, Any]]:
        p = self.project(project_id)
        query = select(PptMessage).where(PptMessage.project_id == p.id)
        if page_id is not None: query = query.where(PptMessage.page_id == page_id)
        rows = self.db.scalars(query.order_by(PptMessage.created_at.desc()).limit(100)).all()
        rows.reverse()
        return [{"id": x.id, "role": x.role, "stage": x.stage, "scope_type": x.scope_type, "page_id": x.page_id, "content_md": x.content_md, "payload": x.structured_payload_json or {}, "created_at": x.created_at.isoformat()} for x in rows]

    def list_checkpoints(self, project_id: str) -> list[dict[str, Any]]:
        p = self.project(project_id)
        rows = self.db.scalars(select(PptCheckpoint).where(PptCheckpoint.project_id == p.id).order_by(PptCheckpoint.created_at.desc())).all()
        return [{"id": x.id, "checkpoint_code": x.checkpoint_code, "stage": x.stage, "status": x.status, "summary_md": x.summary_md, "payload": x.payload_json or {}, "created_at": x.created_at.isoformat()} for x in rows]

    def _message(self, p: PptProject, role: str, content: str, page_id: str | None = None, payload: dict | None = None) -> PptMessage:
        item = PptMessage(project_id=p.id, page_id=page_id, role=role, stage=p.current_stage, scope_type="page" if page_id else "project", content_md=content, structured_payload_json=payload or {})
        self.db.add(item); self.db.flush(); return item

    def confirm_checkpoint(self, project_id: str, code: str, note: str | None = None) -> dict[str, Any]:
        p = self.project(project_id)
        checkpoint = self.db.scalar(select(PptCheckpoint).where(PptCheckpoint.project_id == p.id, PptCheckpoint.checkpoint_code == code, PptCheckpoint.status == "pending").order_by(PptCheckpoint.created_at.desc()))
        if not checkpoint:
            raise HTTPException(404, "确认点不存在或已处理")
        checkpoint.status = "confirmed"
        if note: checkpoint.summary_md = f"{checkpoint.summary_md}\n{note}".strip()
        p.latest_checkpoint_code = None
        # Confirmation is an explicit transition, never a hidden side effect
        # of generating the next artifact.
        if code == "outline_confirm":
            p.current_stage = "visual"
            p.status = "active"
        elif code.endswith("_confirm"):
            p.status = "active"
        self.db.commit()
        return self.serialize_project(p)

    def patch_document(self, project_id: str, page_id: str, document: dict[str, Any], revision: int) -> dict[str, Any]:
        p = self.project(project_id); page = self.page(project_id, page_id)
        if revision != page.document_revision:
            raise HTTPException(409, "页面已被更新，请刷新后重试")
        # Normalize editor payloads at the API boundary so every element has a
        # stable id and bounded geometry before it becomes a new version.
        normalized = dict(document or {})
        normalized["version"] = 2
        normalized["canvas"] = {"width": 1280, "height": 720}
        elements = []
        for index, raw in enumerate(normalized.get("elements") or []):
            item = dict(raw or {})
            item["id"] = str(item.get("id") or f"element-{index + 1}")
            for key in ("x", "y", "w", "h"):
                try:
                    item[key] = max(0.0, min(100.0, float(item.get(key, 0))))
                except (TypeError, ValueError):
                    item[key] = 0.0
            item.setdefault("type", "body")
            item.setdefault("zIndex", index + 1)
            if item.get("type") == "image": item.setdefault("object_fit", "cover")
            elements.append(item)
        normalized["elements"] = elements
        page.document_revision += 1
        page.design_document_json = normalized
        page.statuses_json = {**(page.statuses_json or {}), "design": "ready"}
        version_no = (self.db.scalar(select(PptDocumentVersion).where(PptDocumentVersion.page_id == page.id).order_by(PptDocumentVersion.version_no.desc())) or PptDocumentVersion(version_no=0)).version_no + 1
        version = PptDocumentVersion(project_id=p.id, page_id=page.id, version_no=version_no, document_json=normalized)
        self.db.add(version); self.db.flush(); page.current_document_version_id = version.id
        self._event(p, "workspace.data.updated", {"page_id": page.id, "revision": page.document_revision})
        self.db.commit(); return self.serialize_page(page)

    def restore_document(self, project_id: str, page_id: str, revision: int, direction: str) -> dict[str, Any]:
        p = self.project(project_id); page = self.page(project_id, page_id)
        if revision != page.document_revision:
            raise HTTPException(409, "页面已被更新，请刷新后重试")
        versions = self.db.scalars(select(PptDocumentVersion).where(PptDocumentVersion.page_id == page.id).order_by(PptDocumentVersion.version_no.asc())).all()
        if not versions:
            raise HTTPException(409, "当前页面没有可恢复版本")
        current_no = next((v.version_no for v in versions if v.id == page.current_document_version_id), versions[-1].version_no)
        target_no = current_no - 1 if direction == "undo" else current_no + 1
        target = next((v for v in versions if v.version_no == target_no), None)
        if not target:
            raise HTTPException(409, "没有更多可恢复的版本")
        page.design_document_json = target.document_json
        page.document_revision += 1
        page.current_document_version_id = target.id
        self._event(p, "workspace.data.updated", {"page_id": page.id, "revision": page.document_revision, "direction": direction})
        self.db.commit()
        return self.serialize_page(page)

    def requirements(self, project_id: str) -> dict[str, Any]:
        p = self.project(project_id)
        req = self.db.scalar(select(PptRequirement).where(PptRequirement.project_id == p.id))
        answers = req.answers_json if req else {}
        return {"project_id": p.id, "status": req.status if req else "pending", "page_count_target": p.page_count_target, "answers": answers, "questions": req.questions_json if req else [], "page_count_options": req.page_count_options_json if req else [], "ready_to_outline": bool(req and req.status == "ready"), "suggested_additions": (answers or {}).get("__suggested_additions", []), "brief_summary": (answers or {}).get("__brief_summary", "")}

    def patch_requirements(self, project_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        p = self.project(project_id); req = self.db.scalar(select(PptRequirement).where(PptRequirement.project_id == p.id))
        if not req:
            req = PptRequirement(project_id=p.id); self.db.add(req)
        if payload.get("page_count_target") is not None: p.page_count_target = int(payload["page_count_target"])
        req.answers_json = {**(req.answers_json or {}), **(payload.get("answers") or {})}; self.db.commit()
        return self.requirements(project_id)

    def generate_requirements(self, project_id: str) -> dict[str, Any]:
        p = self.project(project_id); req = self.db.scalar(select(PptRequirement).where(PptRequirement.project_id == p.id))
        gateway = ProviderGateway(self.db, self.user.id)
        data = gateway.json("你是课程课件需求分析师。只输出 JSON：{page_count_options:[{label,page_count,reason}],questions:[{code,label,options:[string]}]}。问题必须帮助教师明确受众、课堂时长和深度。", {"request": p.request_text})
        req.page_count_options_json = data.get("page_count_options", [])[:4]; req.questions_json = data.get("questions", [])[:6]; self.db.commit(); return self.requirements(project_id)

    def _source_context(self, p: PptProject) -> list[dict[str, Any]]:
        rows = self.db.scalars(select(PptSourceDocument).join(PptSourceCollection).where(PptSourceCollection.project_id == p.id)).all()
        fingerprint = hashlib.sha1("|".join(f"{x.id}:{len(x.content_md or '')}:{x.title}" for x in rows).encode()).hexdigest()
        cache_key = f"{p.id}:{fingerprint}"
        with _SOURCE_CONTEXT_LOCK:
            cached = _SOURCE_CONTEXT_CACHE.get(cache_key)
        if cached is not None:
            return cached
        context = [{"id": x.id, "title": x.title, "source_type": x.source_type, "collection_id": x.collection_id, "content": (x.content_md or "")[:6000], "metadata": x.metadata_json or {}} for x in rows]
        with _SOURCE_CONTEXT_LOCK:
            for key in [key for key in _SOURCE_CONTEXT_CACHE if key.startswith(f"{p.id}:") and key != cache_key]:
                _SOURCE_CONTEXT_CACHE.pop(key, None)
            _SOURCE_CONTEXT_CACHE[cache_key] = context
        return context

    def list_sources(self, project_id: str) -> list[dict[str, Any]]:
        p = self.project(project_id)
        rows = self.db.scalars(select(PptSourceDocument).join(PptSourceCollection).where(PptSourceCollection.project_id == p.id).order_by(PptSourceDocument.id.asc())).all()
        return [{"id": x.id, "title": x.title, "source_type": x.source_type, "metadata": x.metadata_json or {}, "collection_id": x.collection_id} for x in rows]

    def _requirement_input(self, project_id: str, content: str | None, option_id: str | None, option_label: str | None, bootstrap: bool):
        p = self.project(project_id)
        req = self.db.scalar(select(PptRequirement).where(PptRequirement.project_id == p.id))
        if not req:
            req = PptRequirement(project_id=p.id); self.db.add(req); self.db.flush()
        text = (option_label or option_id or content or "").strip()
        if not bootstrap and text:
            self._message(p, "user", text, payload={"option_id": option_id, "option_label": option_label} if option_id else {})
        answers = dict(req.answers_json or {})
        transcript_rows = self.db.scalars(select(PptMessage).where(PptMessage.project_id == p.id, PptMessage.page_id.is_(None)).order_by(PptMessage.created_at.desc()).limit(24)).all()
        transcript = [{"role": m.role, "content": m.content_md} for m in reversed(transcript_rows)]
        return p, req, answers, transcript

    def _apply_requirement_data(self, p: PptProject, req: PptRequirement, answers: dict[str, Any], data: dict[str, Any], fallback: dict[str, Any], attachments: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        patch = data.get("answers_patch") if isinstance(data.get("answers_patch"), dict) else {}
        answers.update({str(k): v for k, v in patch.items() if v not in (None, "")})
        for key in ("page_count_target",):
            if data.get(key) not in (None, ""): answers[key] = data[key]
        suggestions = [str(item) for item in (data.get("suggested_additions") or [])[:4] if str(item).strip()]
        answers["__suggested_additions"] = suggestions
        answers["__brief_summary"] = str(data.get("brief_summary") or fallback["brief_summary"])
        req.answers_json = answers
        req.questions_json = [data["question"]] if isinstance(data.get("question"), dict) else (req.questions_json or [])
        req.status = "ready" if bool(data.get("ready_to_outline")) else "collecting"
        if data.get("page_count_target"):
            try: p.page_count_target = int(data["page_count_target"])
            except (TypeError, ValueError): pass
        ready = req.status == "ready"
        assistant_text = str(data.get("assistant_markdown") or fallback["assistant_markdown"])
        if ready and "生成大纲" not in assistant_text: assistant_text += "\n\n需求已经足够明确，你现在可以点击“生成我的大纲”。如果还有页数、案例或风格偏好，也可以继续补充。"
        payload = {"intent_type": "clarify_requirements", "question": data.get("question"), "answers": answers, "missing_fields": data.get("missing_fields") or [], "ready_to_outline": ready, "suggested_additions": suggestions, "brief_summary": answers["__brief_summary"]}
        assistant = self._message(p, "assistant", assistant_text, payload=payload)
        self.db.commit()
        if attachments is None:
            attachments = self.list_sources(p.id)
        return {"project_id": p.id, "message": {"id": assistant.id, "role": assistant.role, "stage": assistant.stage, "scope_type": assistant.scope_type, "content_md": assistant.content_md, "payload": payload, "created_at": assistant.created_at.isoformat()}, "status": req.status, "answers": answers, "question": data.get("question") if isinstance(data.get("question"), dict) else None, "ready_to_outline": ready, "missing_fields": data.get("missing_fields") or [], "suggested_additions": suggestions, "brief_summary": answers["__brief_summary"], "attachments": attachments}

    def requirement_chat(self, project_id: str, content: str | None = None, option_id: str | None = None, option_label: str | None = None, bootstrap: bool = False) -> dict[str, Any]:
        p, req, answers, transcript = self._requirement_input(project_id, content, option_id, option_label, bootstrap)
        prompt = self._requirement_prompt()
        fallback = self._requirement_fallback()
        sources = self._source_context(p)
        try:
            data = ProviderGateway(self.db, self.user.id).json(prompt, {"request": p.request_text, "answers": answers, "history": transcript, "sources": sources})
        except HTTPException:
            data = fallback
        attachments = [{key: source.get(key) for key in ("id", "title", "source_type", "metadata", "collection_id")} for source in sources]
        return self._apply_requirement_data(p, req, answers, data, fallback, attachments)

    @staticmethod
    def _requirement_prompt() -> str:
        return """你是 PPT 需求访谈 Agent。根据项目描述、资料和对话历史，逐轮帮助教师明确课件需求。只输出 JSON：
{"assistant_markdown":"...","brief_summary":"用教师口吻总结当前课件需求（1-3句）","question":{"code":"audience|duration|goals|slide_count|other","label":"...","options":["..."]}|null,"answers_patch":{},"missing_fields":["..."],"ready_to_outline":false,"suggested_additions":["..."],"page_count_target":null}
每轮最多 6 个选项；选项必须是短文本。至少确认受众、教学目标、课堂时长和页数/节奏，不要询问视觉风格，视觉主题由后续步骤人工选择。信息足够时 ready_to_outline=true，并明确告诉教师可以点击生成大纲，但仍给出可补充内容。"""

    @staticmethod
    def _requirement_fallback() -> dict[str, Any]:
        return {"assistant_markdown": "为了做出合适的大纲，还需要确认教学目标和课件节奏。你最希望学生学会什么？", "brief_summary": "正在梳理课件主题、受众与教学节奏。", "question": {"code": "goals", "label": "这套课件最重要的教学目标是什么？", "options": ["理解核心概念", "掌握实践方法", "完成课堂讨论", "准备考试或汇报"]}, "answers_patch": {}, "missing_fields": ["goals"], "ready_to_outline": False, "suggested_additions": [], "page_count_target": None}

    def requirement_chat_stream(self, project_id: str, content: str | None = None, option_id: str | None = None, option_label: str | None = None, bootstrap: bool = False):
        p, req, answers, transcript = self._requirement_input(project_id, content, option_id, option_label, bootstrap)
        prompt, fallback = self._requirement_prompt(), self._requirement_fallback()
        sources = self._source_context(p)
        attachments = [{key: source.get(key) for key in ("id", "title", "source_type", "metadata", "collection_id")} for source in sources]
        try:
            stream = ProviderGateway(self.db, self.user.id).stream_json(prompt, {"request": p.request_text, "answers": answers, "history": transcript, "sources": sources})
            for event in stream:
                if event.get("type") == "chunk":
                    yield event
                elif event.get("type") == "final":
                    yield {"type": "complete", "result": self._apply_requirement_data(p, req, answers, event.get("data") or {}, fallback, attachments)}
        except Exception:
            yield {"type": "complete", "result": self._apply_requirement_data(p, req, answers, fallback, fallback, attachments)}

    def route_message(self, project_id: str, content: str, page_id: str | None = None, option_id: str | None = None, option_label: str | None = None) -> dict[str, Any]:
        p = self.project(project_id)
        if not page_id:
            return self.requirement_chat(project_id, content, option_id, option_label)
        self._message(p, "user", content, page_id)
        decision = ProviderGateway(self.db, self.user.id).json(
            "你是课件工作区 router。只输出 JSON：{action_type,should_execute,reason}. action_type 必须是 page_update_outline_in_search、page_generate_search_queries、page_search_run、page_visual_research、page_visual_plan、page_summary_generate、page_draft_generate、page_design_generate 之一。无法判断时 should_execute=false。",
            {"message": content, "project_stage": p.current_stage, "page_id": page_id},
        )
        if not decision.get("should_execute"):
            assistant = self._message(p, "assistant", decision.get("reason") or "我先把这条需求记下，告诉我希望修改哪一页或先生成哪一步。", page_id, decision)
            self.db.commit()
            return {"decision": decision, "message": {"id": assistant.id, "role": assistant.role, "content_md": assistant.content_md}}
        action = str(decision.get("action_type"))
        if not page_id:
            raise HTTPException(400, "页面级 Agent 动作需要先选择一个页面")
        if action == "page_update_outline_in_search":
            patch = ProviderGateway(self.db, self.user.id).json("只输出 JSON：{title,bullets}，根据用户要求修改当前页面，保留不变的信息。", {"message": content, "current_title": self.page(project_id, page_id).title, "current_bullets": self.page(project_id, page_id).bullets_json})
            result = self.patch_page(project_id, page_id, {"title": patch.get("title") or self.page(project_id, page_id).title, "bullets": patch.get("bullets") or self.page(project_id, page_id).bullets_json})
            self._message(p, "assistant", "已更新当前页结构，研究、初稿和设计稿已标记为待刷新。", page_id, {"action_type": action})
            self.db.commit()
            return {"decision": decision, "result": result}
        result = self.run_action(project_id, page_id, action)
        result["decision"] = decision
        self._message(p, "assistant", "已完成当前页动作。你可以继续修改内容，或让我生成下一阶段产物。", page_id, {"action_type": action, "result": result})
        self.db.commit()
        return result

    def generate_outline(self, project_id: str, page_count_target: int, _legacy_style: str | None = None) -> dict[str, Any]:
        p = self.project(project_id)
        req = self.db.scalar(select(PptRequirement).where(PptRequirement.project_id == p.id))
        if req and req.status != "ready" and p.current_stage == "init":
            raise HTTPException(409, "请先完成需求对话；Agent 确认后才能生成大纲")
        p.page_count_target, p.current_stage = page_count_target, "outline"
        gateway = ProviderGateway(self.db, self.user.id)
        data = gateway.json("你是大学教学课件架构师。只输出 JSON：{sections:[{title,pages:[{title,bullets}]}]}。总页数必须接近目标，页面职责不能重复。", {"request": p.request_text, "page_count": page_count_target})
        pages_payload = [page for section in data.get("sections", []) if isinstance(section, dict) for page in section.get("pages", []) if isinstance(page, dict)]
        if not pages_payload:
            raise HTTPException(502, "模型没有返回有效大纲")
        for old in list(p.pages): self.db.delete(old)
        outline = {"sections": data.get("sections", [])}
        version_no = (self.db.scalar(select(PptOutlineVersion).where(PptOutlineVersion.project_id == p.id).order_by(PptOutlineVersion.version_no.desc())) or PptOutlineVersion(version_no=0)).version_no + 1
        self.db.add(PptOutlineVersion(project_id=p.id, version_no=version_no, outline_json=outline, status="ready"))
        order = 0
        for section in data.get("sections", []):
            for page in section.get("pages", []):
                statuses = {k: ("ready" if k == "outline" else "empty") for k in ("outline", "search", "visual", "summary", "draft", "design")}
                self.db.add(PptPage(project_id=p.id, section_title=section.get("title", ""), sort_order=order, title=page.get("title", "未命名页面"), bullets_json=page.get("bullets", [])[:8], statuses_json=statuses)); order += 1
        p.current_stage = "search"; self.db.commit(); self.db.refresh(p)
        checkpoint = PptCheckpoint(project_id=p.id, checkpoint_code="outline_confirm", stage="outline", status="pending", summary_md="大纲已生成，请确认章节和页数后继续。", payload_json=outline)
        self.db.add(checkpoint); p.latest_checkpoint_code = checkpoint.checkpoint_code
        self._event(p, "outline.completed", {"page_count": order}); self.db.commit()
        return {"project": self.serialize_project(p), "outline": outline, "pages": [self.serialize_page(x) for x in sorted(p.pages, key=lambda x: x.sort_order)]}

    def select_theme(self, project_id: str, theme_id: str) -> dict[str, Any]:
        p = self.project(project_id)
        from app.services.ppt_theme import get_theme
        theme = validate_theme_pack(get_theme(theme_id))
        theme = {**theme, "schema_version": 1, "base_theme": theme.get("base_theme") or theme["id"], "layout_recipes": {layout["id"]: {"roles": layout.get("roles", []), "kind": layout.get("kind")} for layout in list_layouts(theme["id"])}}
        # A theme change invalidates all generated layout/design artifacts.
        p.theme_id = theme["id"]; p.theme_config = theme; p.current_stage = "layout"; p.design_status = "pending"
        p.layout_assignments = {}
        for page in p.pages:
            for version in self.db.scalars(select(PptDocumentVersion).where(PptDocumentVersion.page_id == page.id)).all():
                self.db.delete(version)
            page.design_document_json = None
            page.draft_document_json = None
            page.current_document_version_id = None
            page.document_revision += 1
            page.statuses_json = {**(page.statuses_json or {}), "visual": "stale" if page.visual_plan_json else "empty", "draft": "empty", "design": "empty"}
        self.db.commit()
        return self.serialize_project(p)

    def assign_layouts(self, project_id: str, assignments: dict[str, str] | None = None, mode: str = "manual") -> dict[str, Any]:
        p = self.project(project_id)
        from app.services.ppt_theme import list_layouts
        layouts = list_layouts(p.theme_id or "light-academic")
        valid = {item["id"] for item in layouts}
        current = dict(p.layout_assignments or {})
        if mode == "auto":
            for index, page in enumerate(sorted(p.pages, key=lambda item: item.sort_order)):
                current[page.id] = layouts[index % len(layouts)]["id"]
        else:
            for page_id, layout_id in (assignments or {}).items():
                if self.page(project_id, page_id).id and layout_id in valid:
                    current[page_id] = layout_id
        p.layout_assignments = current
        if p.pages and all(page.id in current for page in p.pages):
            p.current_stage = "design"
        self.db.commit()
        return {"project": self.serialize_project(p), "assignments": current, "layouts": layouts}

    @staticmethod
    def serialize_generation(job: PptGenerationJob) -> dict[str, Any]:
        return {
            "id": job.id, "project_id": job.project_id, "status": job.status,
            "stage": job.stage, "total_pages": job.total_pages,
            "completed_pages": job.completed_pages, "failed_pages": job.failed_pages,
            "current_page_id": job.current_page_id, "error_message": job.error_message,
            "created_at": job.created_at.isoformat(),
            "started_at": job.started_at.isoformat() if job.started_at else None,
            "finished_at": job.finished_at.isoformat() if job.finished_at else None,
            "updated_at": job.updated_at.isoformat(),
        }

    def start_generation(self, project_id: str) -> dict[str, Any]:
        p = self.project(project_id)
        if not p.theme_id:
            raise HTTPException(409, "请先选择主题")
        with _GENERATION_LOCKS_GUARD:
            lock = _GENERATION_LOCKS.setdefault(p.id, threading.Lock())
        with lock:
            existing = self.db.scalar(select(PptGenerationJob).where(
                PptGenerationJob.project_id == p.id,
                PptGenerationJob.status.in_(["queued", "running"]),
            ).order_by(PptGenerationJob.created_at.desc()))
            if existing:
                return self.serialize_generation(existing)
            pages = sorted(p.pages, key=lambda x: x.sort_order)
            p.design_status, p.current_stage = "running", "layout"
            job = PptGenerationJob(project_id=p.id, total_pages=len(pages), status="queued", stage="queued")
            self.db.add(job); self.db.commit(); self.db.refresh(job)
            self._event(p, "generation.started", {"job_id": job.id, "total_pages": len(pages)}); self.db.commit()
            generation_executor.submit(_run_generation_job, job.id, p.id, self.user.id)
            return self.serialize_generation(job)

    def start_visual_research(self, project_id: str) -> dict[str, Any]:
        """Queue visual research immediately and return without waiting on providers."""
        p = self.project(project_id)
        with _GENERATION_LOCKS_GUARD:
            lock = _GENERATION_LOCKS.setdefault(f"visual:{p.id}", threading.Lock())
        with lock:
            existing = self.db.scalar(select(PptGenerationJob).where(
                PptGenerationJob.project_id == p.id,
                PptGenerationJob.stage == "visual",
                PptGenerationJob.status.in_(["queued", "running"]),
            ).order_by(PptGenerationJob.created_at.desc()))
            if existing:
                return self.serialize_generation(existing)
            pages = sorted(p.pages, key=lambda x: x.sort_order)
            job = PptGenerationJob(project_id=p.id, total_pages=len(pages), status="queued", stage="visual")
            self.db.add(job); self.db.commit(); self.db.refresh(job)
            self._event(p, "visual.research.started", {"job_id": job.id, "total_pages": len(pages)}); self.db.commit()
            visual_executor.submit(_run_visual_research_job, job.id, p.id, self.user.id)
            return self.serialize_generation(job)

    def generation_job(self, project_id: str, job_id: str) -> dict[str, Any]:
        self.project(project_id)
        job = self.db.scalar(select(PptGenerationJob).where(PptGenerationJob.id == job_id, PptGenerationJob.project_id == project_id))
        if not job:
            raise HTTPException(404, "生成任务不存在")
        return self.serialize_generation(job)

    def cancel_generation(self, project_id: str, job_id: str) -> dict[str, Any]:
        self.project(project_id)
        job = self.db.scalar(select(PptGenerationJob).where(PptGenerationJob.id == job_id, PptGenerationJob.project_id == project_id))
        if not job: raise HTTPException(404, "生成任务不存在")
        if job.status in {"queued", "running"}:
            job.status, job.stage = "cancelled", "cancelled"; job.finished_at = datetime.now(timezone.utc)
            self._event(self.project(project_id), "generation.cancelled", {"job_id": job.id}); self.db.commit()
        return self.serialize_generation(job)

    def generate_design(self, project_id: str) -> dict[str, Any]:
        p = self.project(project_id)
        if not p.theme_id:
            raise HTTPException(409, "请先选择主题")
        theme = get_theme(p.theme_id); layout_list = list_layouts(p.theme_id); layouts = {item["id"]: item for item in layout_list}
        pages = sorted(p.pages, key=lambda x: x.sort_order)
        p.design_status = "running"; p.current_stage = "layout"; self._event(p, "design.pipeline.started", {"page_count": len(pages)}); self.db.commit()

        # Ask the model for a whole-deck mapping so layout choices are globally coherent.
        outline = [{"id": page.id, "order": page.sort_order, "section": page.section_title, "title": page.title, "bullets": page.bullets_json or []} for page in pages]
        assignments: dict[str, str] = {}
        try:
            data = ProviderGateway(self.db, self.user.id).json(
                "你是教学课件版式总监。只输出 JSON：{assignments:{page_id:layout_id}}。根据完整大纲为每页选择最合适版式。章节开头用 section；比较/流程/内容较多用 two-column；单一金句或结论用 quote；其余用 title-content。避免连续重复，且只能使用提供的合法布局。",
                {"outline": outline, "layouts": layout_list, "theme": theme["colors"]},
            )
            raw = data.get("assignments") if isinstance(data, dict) else {}
            if isinstance(raw, dict):
                valid = set(layouts)
                assignments = {str(page_id): str(layout_id) for page_id, layout_id in raw.items() if str(page_id) in {x.id for x in pages} and str(layout_id) in valid}
        except Exception:
            assignments = {}
        # Deterministic fallback guarantees the user can always reach the editor.
        for index, page in enumerate(pages):
            if page.id in assignments:
                continue
            bullets = page.bullets_json or []
            title = (page.title or "").lower()
            if index == 0 or (page.section_title and (index == 0 or page.section_title != pages[index - 1].section_title)):
                kind = "section"
            elif len(bullets) <= 1 or any(key in title for key in ("总结", "结论", "关键", "takeaway", "quote")):
                kind = "quote"
            elif len(bullets) >= 5 or any(key in title for key in ("比较", "对比", "流程", "步骤", "方法")):
                kind = "two-column"
            else:
                kind = "title-content"
            assignments[page.id] = next((item["id"] for item in layout_list if item["kind"] == kind), layout_list[0]["id"])
        p.layout_assignments = assignments; p.current_stage = "design"; self._event(p, "layout.assigned", {"assignments": assignments}); self.db.commit()

        quality: list[dict[str, Any]] = []
        for page in pages:
            layout = layouts.get(assignments.get(page.id), layouts.get("title-content")) or layout_list[0]
            document: dict[str, Any]
            try:
                document = self._slide_doc(p, page, "design", layout=layout, outline=outline)
            except Exception:
                document = self._template_document(page, theme, layout)
            if not self._document_quality_ok(document):
                try:
                    repaired = self._slide_doc(p, page, "design_repair", layout=layout, outline=outline)
                    document = repaired if self._document_quality_ok(repaired) else self._template_document(page, theme, layout)
                except Exception:
                    document = self._template_document(page, theme, layout)
            page.design_document_json = document
            page.document_revision += 1
            page.statuses_json = {**(page.statuses_json or {}), "design": "ready"}
            version_no = (self.db.scalar(select(PptDocumentVersion).where(PptDocumentVersion.page_id == page.id).order_by(PptDocumentVersion.version_no.desc())) or PptDocumentVersion(version_no=0)).version_no + 1
            version = PptDocumentVersion(project_id=p.id, page_id=page.id, version_no=version_no, document_json=document)
            self.db.add(version); self.db.flush(); page.current_document_version_id = version.id
            quality.append({"page_id": page.id, "layout_id": layout["id"], "elements": len(document.get("elements") or []), "status": "ok"})
            self._event(p, "page.designed", {"page_id": page.id, "layout_id": layout["id"]})
        self._event(p, "design.quality_checked", {"pages": len(quality), "passed": len(quality)}); p.current_stage = "design"; p.design_status = "ready"; self._event(p, "design.completed", {"page_count": len(pages)}); self.db.commit()
        return {"project": self.serialize_project(p), "pages": [self.serialize_page(page) for page in pages], "assignments": assignments, "quality": quality}

    def preview_svg(self, project_id: str, page_id: str, layout_id: str | None = None) -> str:
        p = self.project(project_id); page = self.page(project_id, page_id)
        from app.services.ppt_theme import get_theme, list_layouts, render_slide_svg
        layouts = {item["id"]: item for item in list_layouts(p.theme_id or "light-academic")}
        selected_layout = layout_id or (p.layout_assignments or {}).get(page.id)
        return render_slide_svg(page, get_theme(p.theme_id or "light-academic"), layouts.get(selected_layout))

    def patch_page(self, project_id: str, page_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        p = self.project(project_id); page = self.page(project_id, page_id)
        page.title, page.bullets_json = payload["title"].strip(), payload.get("bullets", [])[:8]
        if payload.get("section_title") is not None: page.section_title = payload["section_title"]
        if payload.get("speaker_notes") is not None: page.speaker_notes = payload["speaker_notes"]
        statuses = dict(page.statuses_json or {}); statuses.update({"search": "stale" if page.citations_json else "empty", "visual": "stale" if page.visual_plan_json else "empty", "summary": "stale" if page.summary_md else "empty", "draft": "stale" if page.draft_document_json else "empty", "design": "stale" if page.design_document_json else "empty"}); page.statuses_json = statuses
        self.db.commit(); return self.serialize_page(page)

    def delete_page(self, project_id: str, page_id: str) -> dict[str, Any]:
        """Delete a slide and compact the storyboard order."""
        p = self.project(project_id)
        page = self.page(project_id, page_id)
        if len(p.pages) <= 1:
            raise HTTPException(409, "课件至少需要保留一页")
        self.db.delete(page)
        self.db.flush()
        remaining_pages = sorted((item for item in p.pages if item.id != page_id), key=lambda x: x.sort_order)
        for index, item in enumerate(remaining_pages):
            item.sort_order = index
        if p.active_page_id == page_id:
            p.active_page_id = remaining_pages[0].id if remaining_pages else None
        p.page_count_target = self.db.query(PptPage).filter(PptPage.project_id == p.id).count()
        p.design_status = "pending"
        self._event(p, "workspace.data.updated", {"deleted_page_id": page_id})
        self.db.commit()
        return {"project": self.serialize_project(p), "deleted_page_id": page_id}

    def insert_page(self, project_id: str, title: str, bullets: list[str], section_title: str = "", document: dict[str, Any] | None = None) -> dict[str, Any]:
        p = self.project(project_id)
        page = PptPage(
            project_id=p.id,
            section_title=section_title.strip(),
            sort_order=len(p.pages),
            title=title.strip() or "未命名页面",
            bullets_json=[str(item) for item in bullets[:8]],
            statuses_json={"outline": "ready", "search": "empty", "visual": "empty", "summary": "empty", "draft": "empty", "design": "ready" if document else "empty"},
            design_document_json=document,
        )
        self.db.add(page)
        self.db.flush()
        p.page_count_target = self.db.query(PptPage).filter(PptPage.project_id == p.id).count()
        p.design_status = "pending"
        if document:
            self.patch_document(project_id, page.id, document, page.document_revision)
        else:
            self.db.commit()
        return self.serialize_page(page)

    def save_visual_selection(self, project_id: str, page_id: str, asset_ids: list[str]) -> dict[str, Any]:
        page = self.page(project_id, page_id)
        ids = list(dict.fromkeys(str(x) for x in asset_ids))
        if len(ids) > 6:
            raise HTTPException(422, "每页最多选择 6 张图片")
        visual = dict(page.visual_plan_json or {})
        assets = [x for x in (visual.get("asset_manifest") or []) if isinstance(x, dict)]
        known = {str(x.get("id")): x for x in assets}
        unknown = [x for x in ids if x not in known]
        if unknown:
            raise HTTPException(422, "选择的图片不属于当前页面候选集")
        visual["selected_asset_ids"] = ids
        visual["selection_status"] = "confirmed"
        page.visual_plan_json = visual
        page.statuses_json = {**(page.statuses_json or {}), "visual": "ready"}
        self.db.commit()
        return self.serialize_page(page)

    def run_action(self, project_id: str, page_id: str | None, action_type: str, replace_existing: bool = True) -> dict[str, Any]:
        p = self.project(project_id); page = self.page(project_id, page_id) if page_id else None
        if action_type == "project_batch_visual":
            return self.start_visual_research(project_id)
        from app.services.ppt_graph import PptWorkflowGraph
        graph_state = PptWorkflowGraph().invoke({"project_id": p.id, "page_id": page_id, "action_type": action_type})
        decision = graph_state["decision"]
        if not decision["should_execute"]:
            raise HTTPException(400, f"Agent 无法执行该动作：{', '.join(decision['missing_data'])}")
        run = PptAgentRun(project_id=p.id, page_id=page.id if page else None, action_type=action_type, status="running", decision_json=decision); self.db.add(run); self.db.flush(); self._event(p, "agent.router.decision", decision); self._event(p, "agent.run.started", {"run_id": run.id, "action_type": action_type, "page_id": page_id}); self.db.commit()
        try:
            if action_type == "page_generate_search_queries": result = self._page_queries(p, page)
            elif action_type in {"page_search_run", "page_search_refresh"}: result = self._page_search(p, page, replace_existing or action_type.endswith("refresh"))
            elif action_type == "page_visual_research": result = self._visual_research(p, page)
            elif action_type == "page_visual_plan": result = self._visual_plan(p, page)
            elif action_type == "page_summary_generate": result = self._summary(p, page)
            elif action_type == "page_draft_generate": result = self._draft(p, page)
            elif action_type == "page_design_generate": result = self._design(p, page)
            elif action_type.startswith("project_batch_"): result = self._batch(p, action_type)
            else: raise HTTPException(400, "不支持的 Agent 动作")
            run.status = "completed"; self._event(p, "agent.run.completed", {"run_id": run.id, "result": result}); self.db.commit(); return {"run_id": run.id, "status": run.status, "result": result}
        except Exception as exc:
            run.status, run.error_message = "failed", str(exc); self._event(p, "agent.run.failed", {"run_id": run.id, "error": str(exc)}); self.db.commit(); raise

    def _page_queries(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        keywords = extract_keywords(getattr(page, "section_title", ""), page.title, getattr(page, "bullets_json", []) or [], getattr(page, "page_role", ""))
        text = query_text(page.section_title, page.title, page.bullets_json or [], page.page_role)
        queries = [{"query_text": text, "query_purpose": "outline keyword match"}]
        for word in keywords[2:6]:
            queries.append({"query_text": f"{word} {page.page_role or 'concept'}", "query_purpose": "keyword expansion"})
        page.search_queries_json = queries[:6]; page.statuses_json = {**(page.statuses_json or {}), "search": "ready"}; self.db.commit(); return {"queries": page.search_queries_json}

    def _page_search(self, p: PptProject, page: PptPage, replace: bool) -> dict[str, Any]:
        if not page.search_queries_json: self._page_queries(p, page)
        if not settings.search_provider_url:
            # DeepSeek's native web_search is exposed by the Responses API.
            # This keeps page search functional without a second search key.
            results: list[dict[str, Any]] = []
            gateway = ProviderGateway(self.db, self.user.id)
            failures = 0
            def search_one(query):
                try:
                    data = gateway.web_search_json(str(query.get("query_text") or ""))
                    items = data.get("results") or data.get("images") or []
                    return [item for item in items if isinstance(item, dict)] if isinstance(items, list) else []
                except Exception:
                    return None
            with ThreadPoolExecutor(max_workers=3) as pool:
                for found in pool.map(search_one, page.search_queries_json[:6]):
                    if found is None: failures += 1
                    else: results.extend(found)
            page.citations_json = self._dedupe_search_results(results)[:30]
            search_status = "failed" if failures and not page.citations_json else "ready" if page.citations_json else "empty"
            page.statuses_json = {**(page.statuses_json or {}), "search": search_status, "visual": "stale"}
            self.db.commit()
            return {"result_count": len(page.citations_json), "failed_queries": failures, "status": search_status, "provider": "deepseek_web_search"}
        import httpx
        results = []
        failures = 0
        headers = {"Authorization": f"Bearer {settings.search_provider_key}"} if settings.search_provider_key else {}
        for q in page.search_queries_json:
            for attempt in range(2):
                try:
                    request_payload = {"query": q["query_text"], "type": "web_and_image"} if attempt == 0 else {"query": q["query_text"]}
                    response = httpx.post(settings.search_provider_url, json=request_payload, headers=headers, timeout=60)
                    response.raise_for_status(); payload = response.json()
                    if isinstance(payload, list):
                        results.extend(payload)
                    else:
                        results.extend(payload.get("results", [])); results.extend(payload.get("images", []))
                    break
                except Exception:
                    if attempt == 1:
                        failures += 1
                        continue
        page.citations_json = self._dedupe_search_results(results)[:30]
        search_status = "failed" if failures and not page.citations_json else "ready"
        page.statuses_json = {**(page.statuses_json or {}), "search": search_status, "visual": "stale"}
        self.db.commit(); return {"result_count": len(page.citations_json), "failed_queries": failures, "status": search_status}

    @staticmethod
    def _wikimedia_image_search(query: str, limit: int = 8) -> list[dict[str, Any]]:
        """Use Wikimedia Commons as a keyless, license-aware image fallback.

        A small public fallback keeps the visual step useful when native search
        cannot return image URLs. A configured provider remains the preferred
        source for richer web/image retrieval.
        """
        if not query.strip():
            return []
        try:
            import httpx

            response = httpx.get(
                "https://commons.wikimedia.org/w/api.php",
                params={
                    "action": "query",
                    "generator": "search",
                    "gsrsearch": query,
                    "gsrnamespace": 6,
                    "gsrlimit": max(1, min(limit, 12)),
                    "prop": "imageinfo",
                    "iiprop": "url|mime|size|extmetadata",
                    "iiurlwidth": 1600,
                    "format": "json",
                    "formatversion": 2,
                },
                headers={"User-Agent": "HKU-Courseware-Agent/1.0"},
                timeout=4,
            )
            response.raise_for_status()
            pages = ((response.json() or {}).get("query") or {}).get("pages") or []
            results: list[dict[str, Any]] = []
            for item in pages:
                info = (item.get("imageinfo") or [{}])[0]
                image_url = info.get("thumburl") or info.get("url")
                if not image_url:
                    continue
                metadata = info.get("extmetadata") or {}

                def meta(name: str) -> str:
                    value = metadata.get(name) or {}
                    return str(value.get("value") or "").strip() if isinstance(value, dict) else str(value).strip()

                title = str(item.get("title") or "").removeprefix("File:").strip()
                results.append({
                    "title": title or "Wikimedia Commons image",
                    "image_url": image_url,
                    "url": info.get("descriptionurl") or image_url,
                    "snippet": meta("ImageDescription") or title,
                    "license": meta("LicenseShortName") or "Wikimedia Commons",
                    "score": 0.62,
                    "source": "wikimedia_commons",
                })
            return results
        except Exception:
            return []

    @staticmethod
    def _openverse_image_search(query: str, limit: int = 8) -> list[dict[str, Any]]:
        """Public, attribution-aware image search fallback."""
        if not query.strip():
            return []
        try:
            import httpx
            response = httpx.get("https://api.openverse.org/v1/images/", params={"q": query, "page_size": max(1, min(limit, 20))}, headers={"User-Agent": "HKU-Courseware-Agent/1.0"}, timeout=6)
            response.raise_for_status()
            results = []
            for item in (response.json() or {}).get("results", []):
                image_url = item.get("thumbnail") or item.get("url")
                if not image_url:
                    continue
                results.append({"title": item.get("title") or "Openverse image", "image_url": image_url, "url": item.get("foreign_landing_url") or item.get("url"), "snippet": item.get("description") or item.get("title") or query, "license": item.get("license") or "Openverse", "score": float(item.get("score") or 0.55), "source": "openverse"})
            return results
        except Exception:
            return []

    @staticmethod
    def _dedupe_search_results(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
        seen: set[str] = set(); output: list[dict[str, Any]] = []
        for raw in results:
            if not isinstance(raw, dict):
                continue
            item = dict(raw)
            key = str(item.get("url") or item.get("link") or item.get("title") or "").strip().lower()
            if not key or key in seen:
                continue
            seen.add(key); output.append(item)
        return output

    @staticmethod
    def _image_url(item: dict[str, Any]) -> str:
        for key in ("image_url", "thumbnail", "thumbnail_url", "image", "src", "url"):
            value = item.get(key)
            if isinstance(value, dict):
                value = value.get("url") or value.get("src")
            if isinstance(value, str) and value.startswith(("http://", "https://", "data:")):
                if key == "url" and not re.search(r"\.(png|jpe?g|webp|gif)(\?|$)", value, re.I):
                    continue
                return value
        return ""

    def _download_image(self, p: PptProject, page: PptPage, url: str, index: int) -> dict[str, Any] | None:
        if not url or url.startswith("data:"):
            return {"id": f"asset-{page.id}-{index}", "src": url, "public_url": url, "source_url": url, "status": "ready"} if url else None
        try:
            parsed_url = urlparse(url)
            if parsed_url.scheme not in {"http", "https"} or not parsed_url.hostname:
                return None
            host = parsed_url.hostname.lower()
            if host in {"localhost", "127.0.0.1", "::1"}:
                return None
            try:
                resolved = socket.gethostbyname(host)
                if ipaddress.ip_address(resolved).is_private or ipaddress.ip_address(resolved).is_loopback or ipaddress.ip_address(resolved).is_link_local:
                    return None
            except (OSError, ValueError):
                return None
            import httpx
            response = httpx.get(url, follow_redirects=True, timeout=6, headers={"User-Agent": "HKU-Courseware-Agent/1.0"})
            response.raise_for_status()
            content = response.content
            if len(content) > 8 * 1024 * 1024:
                return None
            content_type = response.headers.get("content-type", "")
            suffix = Path(urlparse(str(response.url)).path).suffix.lower()
            if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
                suffix = ".jpg" if "jpeg" in content_type or "jpg" in content_type else ".png"
            root = settings.ppt_storage_path / p.id / "assets"
            root.mkdir(parents=True, exist_ok=True)
            filename = f"{stable_asset_id(url)}{suffix}"
            path = root / filename
            if not path.exists():
                path.write_bytes(content)
            public_url = f"/api/ppt/projects/{p.id}/assets/{filename}"
            asset = {"id": f"asset-{page.id}-{index}", "src": public_url, "asset_path": str(path), "public_url": public_url, "source_url": url, "mime_type": content_type or mimetypes.guess_type(filename)[0] or "image/jpeg", "status": "ready"}
            # Recognition is deliberately deferred; it must never block image insertion.
            asset["recognition"] = {"status": "pending"}
            return asset
        except Exception as exc:
            return {"id": f"asset-{page.id}-{index}", "source_url": url, "status": "failed", "error": str(exc)[:160]}

    def _recognize_image(self, content: bytes, mime_type: str) -> dict[str, Any]:
        """Best-effort OCR/vision enrichment; never blocks slide generation."""
        try:
            gateway = ProviderGateway(self.db, self.user.id)
            client = gateway._client()
            model = gateway.config.model if gateway.config else settings.llm_model
            response = client.chat.completions.create(model=model, temperature=0.1, messages=[{"role": "user", "content": [{"type": "text", "text": "请识别这张图片：输出 JSON，字段为 description、ocr_text、visual_tags、is_chart。不要猜测无法看清的文字。"}, {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{base64.b64encode(content).decode()}"}}]}])
            text = response.choices[0].message.content or "{}"
            try:
                parsed = json.loads(text)
            except json.JSONDecodeError:
                match = re.search(r"\{[\s\S]*\}", text)
                parsed = json.loads(match.group(0)) if match else {"description": text}
            return {"description": str(parsed.get("description") or ""), "ocr_text": str(parsed.get("ocr_text") or ""), "visual_tags": [str(item) for item in (parsed.get("visual_tags") or [])[:8]], "is_chart": bool(parsed.get("is_chart"))}
        except Exception as exc:
            return {"status": "unavailable", "error": str(exc)[:120]}

    def _visual_research(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        if not getattr(page, "search_queries_json", None) and not getattr(page, "citations_json", None):
            self._page_queries(p, page)
        keywords = extract_keywords(getattr(page, "section_title", ""), page.title, getattr(page, "bullets_json", []) or [], getattr(page, "page_role", ""))
        query_text_value = build_query(getattr(page, "section_title", ""), page.title, getattr(page, "bullets_json", []) or [], getattr(page, "page_role", ""))
        raw_results: list[dict[str, Any]] = []
        existing_citations = getattr(page, "citations_json", None) or []
        # Existing citations with direct image URLs are already usable. Avoid
        # replacing a deterministic candidate set with a second network search
        # (and keep retries from duplicating or reordering the same assets).
        has_direct_image = any(isinstance(item, dict) and self._image_url(item) for item in existing_citations)
        if has_direct_image:
            browser_error = ""
            raw_results.extend(existing_citations)
        else:
            try:
                raw_results.extend(BrowserImageSearch(mode=settings.visual_search_mode, proxy=settings.visual_search_proxy or None, timeout=settings.visual_search_timeout_seconds, headless=settings.visual_search_browser_headless, executable_path=settings.visual_search_browser_executable_path or None).search_sync(query_text_value, 12))
            except Exception as exc:
                browser_error = str(exc)[:200]
            else:
                browser_error = ""
            raw_results.extend(existing_citations)
        if not raw_results:
            for query in (page.search_queries_json or [])[:3]:
                text = str(query.get("query_text") or "")
                found = self._wikimedia_image_search(text, 8)
                raw_results.extend(found or self._openverse_image_search(text, 8))
        ranked = rank_candidates(raw_results, keywords, limit=12)
        page.citations_json = self._dedupe_search_results(ranked)
        candidates: list[dict[str, Any]] = []
        for item in page.citations_json or []:
            if not isinstance(item, dict):
                continue
            url = self._image_url(item)
            if not url:
                continue
            if url.startswith("data:"):
                asset = self._download_image(p, page, url, len(candidates) + 1)
            else:
                try:
                    cached = download_image(url, settings.ppt_storage_path / p.id / "assets", f"/api/ppt/projects/{p.id}/assets", max_bytes=settings.visual_search_max_image_bytes, timeout=settings.visual_search_timeout_seconds)
                    asset = {"id": f"asset-{page.id}-{len(candidates)+1}", "source_url": item.get("url") or url, **cached}
                except Exception:
                    # A search result can be a perfectly valid image even when
                    # the backend cannot cache it (hotlink protection, a
                    # transient CDN error, or a content-type mismatch). Keep a
                    # public URL as a degraded candidate so the browser can
                    # still render it and the user can continue with visuals.
                    asset = self._remote_image_asset(url, page, len(candidates) + 1)
            if asset and asset.get("status") == "failed":
                asset = self._remote_image_asset(url, page, len(candidates) + 1)
            if asset:
                asset.update({"title": item.get("title") or page.title, "alt": item.get("snippet") or item.get("title") or page.title, "score": float(item.get("score") or item.get("relevance") or 0.5), "license": item.get("license") or item.get("copyright")})
                candidates.append(asset)
            if len(candidates) >= 6:
                break
        if not candidates:
            # Existing citations may contain web pages but no direct image URL.
            # Keep the image step useful by querying the public image fallback
            # before marking the page as empty.
            queries = page.search_queries_json or [{"query_text": f"{page.title} {page.section_title}".strip()}]
            for query in queries[:1]:
                text = str(query.get("query_text") or "")
                fallback = self._wikimedia_image_search(text, 4) or self._openverse_image_search(text, 4)
                for item in fallback:
                    url = self._image_url(item)
                    asset = self._download_image(p, page, url, len(candidates) + 1) if url else None
                    if asset and asset.get("status") == "failed":
                        asset = self._remote_image_asset(url, page, len(candidates) + 1)
                    if asset:
                        asset.update({"title": item.get("title") or page.title, "alt": item.get("snippet") or item.get("title") or page.title, "score": float(item.get("score") or 0.5), "license": item.get("license")})
                        candidates.append(asset)
                    if len(candidates) >= 6:
                        break
                if len(candidates) >= 6:
                    break
        visual = dict(page.visual_plan_json or {})
        # A new candidate set requires an explicit user confirmation again.
        if hasattr(page, "search_queries_json"):
            visual["selected_asset_ids"] = []
            visual["selection_status"] = "pending"
        candidates = rerank(candidates, query_text(getattr(page, "section_title", ""), page.title, getattr(page, "bullets_json", []) or [], getattr(page, "page_role", "")), getattr(settings, "clip_model_name", ""))
        visual["asset_manifest"] = candidates[:6]
        visual["asset_query"] = query_text(getattr(page, "section_title", ""), page.title, getattr(page, "bullets_json", []) or [], getattr(page, "page_role", ""))
        visual["asset_status"] = "ready" if candidates else ("failed" if browser_error else "empty")
        if browser_error and not candidates: visual["asset_error"] = browser_error
        page.visual_plan_json = visual
        page.statuses_json = {**(page.statuses_json or {}), "visual": "ready" if candidates else "empty"}
        self.db.commit()
        return {"asset_count": len(candidates), "assets": candidates, "status": visual["asset_status"]}

    @staticmethod
    def _remote_image_asset(url: str, page: PptPage, index: int) -> dict[str, Any] | None:
        """Return a browser-renderable fallback after local caching fails.

        Only public HTTP(S) hosts are allowed here; this mirrors the SSRF
        guard used by ``_download_image`` and prevents a failed cache request
        from turning into an unsafe remote URL in the client.
        """
        try:
            validate_public_url(url)
        except ValueError:
            return None
        return {
            "id": f"asset-{page.id}-{index}",
            "src": url,
            "public_url": url,
            "source_url": url,
            "status": "ready",
            "remote_only": True,
        }

    def _visual_plan(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        visual = dict(page.visual_plan_json or {})
        all_assets = [item for item in visual.get("asset_manifest", []) if isinstance(item, dict) and item.get("status") == "ready"]
        selected = visual.get("selected_asset_ids")
        assets = [item for item in all_assets if not (isinstance(selected, list) and item.get("id") not in selected)]
        role = page.page_role or "concept"
        fallback_kind = {"cover": "image-focus", "section": "image-focus", "case": "image-text", "process": "process-timeline", "comparison": "two-column"}.get(role, "image-text" if assets else "content")
        plan = {"layout_kind": fallback_kind, "image_slots": []}
        for index, asset in enumerate(assets[:6]):
            if index == 0 and role in {"cover", "section"}:
                x, y, w, h, hint = 0, 0, 100, 100, "background"
            elif len(assets) == 1:
                x, y, w, h, hint = 54, 18, 40, 70, "right"
            elif len(assets) <= 3:
                x, y, w, h, hint = 54 + (index % 2) * 21, 22 + (index // 2) * 34, 19, 28, "supporting"
            else:
                x, y, w, h, hint = 54 + (index % 3) * 15, 18 + (index // 3) * 34, 13, 28, "grid"
            plan["image_slots"].append({"slot_id": f"image-{index + 1}", "asset_id": asset["id"], "src": asset.get("public_url") or asset.get("src"), "asset_path": asset.get("asset_path"), "placement_hint": hint, "x": x, "y": y, "w": w, "h": h, "object_fit": "cover", "confidence": round(min(0.98, 0.62 + float(asset.get("score") or 0) * 0.3), 2)})
        visual.update(plan)
        page.visual_plan_json = visual
        page.statuses_json = {**(page.statuses_json or {}), "visual": "ready"}
        self.db.commit()
        return visual

    def _summary(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        text = "\n\n".join(c.content_md for col in self.db.scalars(select(PptSourceCollection).where(PptSourceCollection.project_id == p.id, PptSourceCollection.page_id == page.id)) for d in col.documents for c in d.chunks)
        if not text and not page.citations_json: raise HTTPException(409, "当前页没有资料，无法生成 Summary")
        gateway = ProviderGateway(self.db, self.user.id); page.summary_md = gateway.text("根据给定资料生成有引用意识的教学 Summary，使用 Markdown，禁止编造事实。", {"title": page.title, "bullets": page.bullets_json, "sources": text or page.citations_json}); page.statuses_json = {**(page.statuses_json or {}), "summary": "ready"}; self.db.commit(); return {"summary_length": len(page.summary_md)}

    def _template_document(self, page: PptPage, theme: dict[str, Any], layout: dict[str, Any]) -> dict[str, Any]:
        colors = theme.get("colors") or {}
        bullets = [str(item) for item in (page.bullets_json or [])[:6]]
        kind = layout.get("kind", "content")
        role = page.page_role or "concept"
        title_size = 42 if role in {"cover", "section"} else 32
        elements: list[dict[str, Any]] = [{"id": "title", "type": "title", "x": 7, "y": 8, "w": 86, "h": 14, "text": page.title, "font_size": title_size, "font_weight": 800, "color": colors.get("title")}]
        if role == "cover":
            elements.append({"id": "cover-subtitle", "type": "caption", "x": 9, "y": 38, "w": 76, "h": 12, "text": page.section_title or "课程课件", "font_size": 22, "font_weight": 500, "color": colors.get("body")})
        elif role == "catalogue":
            for index, text in enumerate(bullets[:8]):
                elements.append({"id": f"catalogue-{index}", "type": "body", "x": 10, "y": 29 + index * 7, "w": 78, "h": 6, "text": f"{index + 1:02d}  {text}", "font_size": 20, "font_weight": 600 if index < 3 else 500, "color": colors.get("body")})
        elif role == "section" or kind == "section":
            elements.append({"id": "section", "type": "body", "x": 10, "y": 43, "w": 80, "h": 16, "text": page.section_title or page.title, "font_size": 34, "font_weight": 700, "color": colors.get("title")})
        elif kind == "quote":
            elements.append({"id": "quote", "type": "body", "x": 12, "y": 38, "w": 76, "h": 24, "text": bullets[0] if bullets else page.title, "font_size": 26, "color": colors.get("body")})
        elif kind == "columns":
            midpoint = max(1, (len(bullets) + 1) // 2)
            for index, text in enumerate(bullets):
                column, row = (0, index) if index < midpoint else (1, index - midpoint)
                elements.append({"id": f"bullet-{index}", "type": "body", "x": 9 + column * 44, "y": 29 + row * 12, "w": 38, "h": 9, "text": f"• {text}", "font_size": 18, "font_weight": 500, "color": colors.get("body")})
        else:
            for index, text in enumerate(bullets):
                elements.append({"id": f"bullet-{index}", "type": "body", "x": 9, "y": 29 + index * 10, "w": 82, "h": 8, "text": f"• {text}", "font_size": 18 if len(text) < 70 else 15, "font_weight": 500, "color": colors.get("body")})
        visual = page.visual_plan_json or {}
        for index, slot in enumerate(visual.get("image_slots") or []):
            if slot.get("src"):
                elements.append({"id": f"image-slot-{index + 1}", "type": "image", "asset_id": slot.get("asset_id"), "src": slot.get("src"), "asset_path": slot.get("asset_path"), "alt": page.title, "x": slot.get("x", 54), "y": slot.get("y", 18), "w": slot.get("w", 40), "h": slot.get("h", 70), "object_fit": slot.get("object_fit", "cover"), "zIndex": 0})
        editor_theme = {**colors, "background": colors.get("bg")}
        return {"version": 2, "canvas": {"width": 1280, "height": 720}, "layout": layout["id"], "theme": editor_theme, "elements": elements, "speaker_notes": page.speaker_notes}

    @staticmethod
    def _document_quality_ok(document: dict[str, Any]) -> bool:
        elements = document.get("elements") if isinstance(document, dict) else None
        if not isinstance(elements, list) or not elements:
            return False
        title_sizes = [float(item.get("font_size", 0)) for item in elements if isinstance(item, dict) and str(item.get("type", "")).lower() in {"title", "heading"}]
        max_title = max(title_sizes, default=0)
        for element in elements:
            if not isinstance(element, dict) or not str(element.get("text") or element.get("src") or "").strip():
                continue
            try:
                x, y, w, h = (float(element.get(key, 0)) for key in ("x", "y", "w", "h"))
            except (TypeError, ValueError):
                return False
            if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > 100.5 or y + h > 100.5:
                return False
            text = str(element.get("text") or "")
            size = float(element.get("font_size", 18) or 18)
            if text and len(text) > 220 and h < 12:
                return False
            if str(element.get("type", "")).lower() in {"body", "caption"} and max_title and size > max_title:
                return False
        return True

    def _slide_doc(self, p: PptProject, page: PptPage, mode: str, layout: dict[str, Any] | None = None, outline: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        gateway = ProviderGateway(self.db, self.user.id)
        selected_layout = layout or {"id": (page.document or {}).get("layout", "title-content"), "kind": "content"}
        raw = gateway.json("你是资深教学课件设计 Agent。只输出结构化 Slide JSON：{layout,theme,elements:[{type,x,y,w,h,text,items,src,font_size,font_weight,color,fill}],speaker_notes}。大纲只是页面职责，不是最终全文；必须优先使用 content_plan 扩写后的论点、案例和视觉建议。根据页面角色设计清晰层级：标题 30-44px、正文 18-24px、辅助文字 12-16px；封面/目录/章节/对比/流程/案例/总结使用匹配的视觉结构。所有坐标使用 0-100 百分比，留出安全边距，禁止严重重叠和文字堆叠。", {"mode": mode, "page": {"title": page.title, "section": page.section_title, "role": page.page_role, "bullets": page.bullets_json, "summary": page.summary_md, "content_plan": page.content_plan_json or {}, "visual_plan": page.visual_plan_json or {}}, "outline": outline or [], "layout": selected_layout, "theme": p.theme_config or {}})
        return self._apply_visual_plan(page, self._normalize_document(raw, selected_layout, p.theme_config or {}))

    @staticmethod
    def _apply_visual_plan(page: PptPage, document: dict[str, Any]) -> dict[str, Any]:
        result = dict(document or {})
        elements = [dict(item) for item in result.get("elements") or [] if isinstance(item, dict)]
        existing_asset_ids = {str(item.get("asset_id")) for item in elements if item.get("asset_id")}
        plan = page.visual_plan_json or {}
        for index, slot in enumerate(plan.get("image_slots") or []):
            if not isinstance(slot, dict) or not slot.get("src") or str(slot.get("asset_id")) in existing_asset_ids:
                continue
            elements.append({"id": f"image-slot-{index + 1}", "type": "image", "asset_id": slot.get("asset_id"), "src": slot.get("src"), "asset_path": slot.get("asset_path"), "alt": slot.get("alt") or page.title, "x": slot.get("x", 54), "y": slot.get("y", 18), "w": slot.get("w", 40), "h": slot.get("h", 70), "object_fit": slot.get("object_fit", "cover"), "zIndex": 0})
        result["elements"] = elements
        result["visual_plan"] = plan
        return PptAgentService._repair_document_layout(result)

    @staticmethod
    def _repair_document_layout(document: dict[str, Any]) -> dict[str, Any]:
        """Apply conservative geometry repairs after model/layout composition."""
        elements = [dict(item) for item in document.get("elements") or [] if isinstance(item, dict)]
        images = [item for item in elements if str(item.get("type") or "").lower() == "image" and float(item.get("x", 0) or 0) > 0]
        text_items = [item for item in elements if str(item.get("type") or "").lower() in {"body", "caption"}]
        for image in images:
            ix, iy, iw, ih = (float(image.get(key, 0) or 0) for key in ("x", "y", "w", "h"))
            for text in text_items:
                tx, ty, tw, th = (float(text.get(key, 0) or 0) for key in ("x", "y", "w", "h"))
                intersects = ix < tx + tw and ix + iw > tx and iy < ty + th and iy + ih > ty
                if not intersects:
                    continue
                if ix >= 50:
                    text["w"] = max(8, min(tw, ix - tx - 3))
                else:
                    image["x"] = min(96 - iw, max(ix, tx + tw + 3))
        for item in elements:
            for key in ("x", "y", "w", "h"):
                try:
                    item[key] = max(0.0, min(100.0, float(item.get(key, 0))))
                except (TypeError, ValueError):
                    item[key] = 0.0
            item["w"] = min(item["w"], 100 - item["x"]); item["h"] = min(item["h"], 100 - item["y"])
        return {**document, "elements": elements}

    @staticmethod
    def _normalize_document(raw: dict[str, Any], selected_layout: dict[str, Any], theme_config: dict[str, Any]) -> dict[str, Any]:
        normalized = dict(raw or {}); normalized["version"] = 2; normalized["canvas"] = {"width": 1280, "height": 720}; normalized["elements"] = []
        normalized["layout"] = selected_layout["id"]; colors = (theme_config or {}).get("colors", {}); normalized["theme"] = {**colors, "background": colors.get("bg"), "tokens": (theme_config or {}).get("tokens", {}), "accessibility": (theme_config or {}).get("accessibility", {})}
        for index, item in enumerate(raw.get("elements") or []):
            element = dict(item or {}); element["id"] = str(element.get("id") or f"element-{index + 1}")
            for key in ("x", "y", "w", "h"):
                try: element[key] = max(0.0, min(100.0, float(element.get(key, 0))))
                except (TypeError, ValueError): element[key] = 0.0
            element.setdefault("type", "body")
            element.setdefault("zIndex", index + 1)
            if element.get("type") == "image": element.setdefault("object_fit", "cover")
            normalized["elements"].append(element)
        return normalized

    def _draft(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        if not page.summary_md: raise HTTPException(409, "请先生成当前页 Summary")
        page.draft_document_json = self._slide_doc(p, page, "draft"); page.statuses_json = {**(page.statuses_json or {}), "draft": "ready"}; self.db.commit(); return {"page_id": page.id, "document": page.draft_document_json}

    def _design(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        if not page.draft_document_json: raise HTTPException(409, "请先生成当前页 Draft")
        page.design_document_json = self._slide_doc(p, page, "design"); page.document_revision += 1
        version_no = (self.db.scalar(select(PptDocumentVersion).where(PptDocumentVersion.page_id == page.id).order_by(PptDocumentVersion.version_no.desc())) or PptDocumentVersion(version_no=0)).version_no + 1
        version = PptDocumentVersion(project_id=p.id, page_id=page.id, version_no=version_no, document_json=page.design_document_json)
        self.db.add(version); self.db.flush(); page.current_document_version_id = version.id
        page.statuses_json = {**(page.statuses_json or {}), "design": "ready"}; self.db.commit(); return {"page_id": page.id, "document": page.design_document_json, "revision": page.document_revision}

    def _batch(self, p: PptProject, action: str) -> dict[str, Any]:
        mapping = {"project_batch_search": "page_search_run", "project_batch_visual": "page_visual_research", "project_batch_summary": "page_summary_generate", "project_batch_draft": "page_draft_generate", "project_batch_design": "page_design_generate"}; results = []
        for page in sorted(p.pages, key=lambda x: x.sort_order):
            try: results.append(self.run_action(p.id, page.id, mapping[action])["result"])
            except HTTPException as exc: results.append({"page_id": page.id, "status": "skipped", "reason": exc.detail})
        return {"items": results}

    def upload_document(self, project_id: str, page_id: str | None, upload: UploadFile) -> dict[str, Any]:
        p = self.project(project_id); page = self.page(project_id, page_id) if page_id else None
        name = Path(upload.filename or "document").name
        suffix = Path(name).suffix.lower()
        allowed = {".pdf", ".md", ".markdown", ".txt", ".csv", ".png", ".jpg", ".jpeg", ".webp", ".pptx", ".docx"}
        if suffix not in allowed: raise HTTPException(415, "仅支持 PDF、Markdown、图片、PPTX、DOCX、TXT、CSV")
        content = upload.file.read()
        if len(content) > 25 * 1024 * 1024: raise HTTPException(413, "单个资料不能超过 25MB")
        root = settings.ppt_storage_path / p.id; root.mkdir(parents=True, exist_ok=True); path = root / name; path.write_bytes(content)
        text = self._extract_text(path, content); collection = self.db.scalar(select(PptSourceCollection).where(PptSourceCollection.project_id == p.id, PptSourceCollection.page_id == (page.id if page else None)))
        if not collection: collection = PptSourceCollection(project_id=p.id, page_id=page.id if page else None, collection_type="page_corpus" if page else "init_corpus", title="页面资料池" if page else "项目资料池"); self.db.add(collection); self.db.flush()
        doc = PptSourceDocument(collection_id=collection.id, source_type="upload", source_uri=str(path), title=name, content_md=text, metadata_json={"mime_type": upload.content_type}); self.db.add(doc); self.db.flush()
        chunks = self._chunk_text(text); doc.chunks = [PptSourceChunk(chunk_index=i, content_md=chunk) for i, chunk in enumerate(chunks)]; self.db.commit(); return {"document_id": doc.id, "title": name, "chunk_count": len(chunks), "collection_type": collection.collection_type}

    def _chunk_text(self, text: str) -> list[str]:
        try:
            from langchain_text_splitters import RecursiveCharacterTextSplitter
            return RecursiveCharacterTextSplitter(chunk_size=1800, chunk_overlap=200).split_text(text) or [""]
        except ImportError:
            return [text[i:i + 1800] for i in range(0, len(text), 1600)] or [""]

    def _extract_text(self, path: Path, content: bytes) -> str:
        if path.suffix.lower() in {".txt", ".md", ".markdown", ".csv"}: return content.decode("utf-8", errors="ignore")
        if path.suffix.lower() == ".pdf":
            try:
                import pypdf
                return "\n".join((page.extract_text() or "") for page in pypdf.PdfReader(str(path)).pages)
            except ImportError: raise HTTPException(500, "解析 PDF 需要 pypdf")
        if path.suffix.lower() == ".docx":
            try:
                from docx import Document
                return "\n".join(p.text for p in Document(str(path)).paragraphs)
            except ImportError: raise HTTPException(500, "解析 DOCX 需要 python-docx")
        if path.suffix.lower() == ".pptx":
            try:
                from pptx import Presentation
                return "\n".join(shape.text for slide in Presentation(str(path)).slides for shape in slide.shapes if hasattr(shape, "text"))
            except ImportError: raise HTTPException(500, "解析 PPTX 需要 python-pptx")
        if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
            # Vision is best-effort: the original image remains available even
            # when the configured provider does not support image inputs.
            try:
                import base64
                gateway = ProviderGateway(self.db, self.user.id)
                client = gateway._client()
                model = gateway.config.model if gateway.config else settings.llm_model
                mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}[path.suffix.lower()]
                response = client.chat.completions.create(model=model, temperature=0.1, messages=[{"role": "user", "content": [{"type": "text", "text": "请提取图片中的文字，并描述图表、关键视觉信息和可用于教学课件的要点。"}, {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{base64.b64encode(content).decode()}"}}]}])
                return response.choices[0].message.content or "图片资料（未提取到文字）"
            except Exception:
                return "图片资料已上传；当前模型不支持视觉解析，请教师在对话中补充图片要点。"
        raise HTTPException(415, "仅支持 PDF、PPTX、DOCX、TXT、MD、CSV")

    def _event(self, p: PptProject, event_type: str, payload: dict[str, Any]) -> None:
        self.db.add(PptAgentEvent(project_id=p.id, event_type=event_type, payload_json=payload))


def _run_visual_research_job(job_id: str, project_id: str, user_id: str) -> None:
    db = SessionLocal()
    try:
        job, project, user = db.get(PptGenerationJob, job_id), db.get(PptProject, project_id), db.get(User, user_id)
        if not job or not project or not user: return
        job.status = "running"; db.commit()
        pages = sorted(project.pages, key=lambda x: x.sort_order)
        def one(page_id: str):
            local = SessionLocal()
            try:
                p, u, page = local.get(PptProject, project_id), local.get(User, user_id), local.get(PptPage, page_id)
                if not p or not u or not page: return False
                svc = PptAgentService(local, u)
                page.statuses_json = {**(page.statuses_json or {}), "visual": "running"}; local.commit()
                svc._event(p, "page.visual.searching", {"job_id": job_id, "page_id": page_id}); local.commit()
                svc._visual_research(p, page); svc._visual_plan(p, page)
                svc._event(p, "page.visual.ready", {"job_id": job_id, "page_id": page_id, "asset_count": len((page.visual_plan_json or {}).get("asset_manifest") or [])}); local.commit()
                return True
            except Exception as exc:
                local.rollback(); page = local.get(PptPage, page_id)
                if page:
                    page.statuses_json = {**(page.statuses_json or {}), "visual": "failed"}
                    page.visual_plan_json = {**(page.visual_plan_json or {}), "asset_status": "failed", "asset_error": str(exc)[:200]}
                    local.commit()
                return False
            finally: local.close()
        # Keep pages in a project serial to reduce search-engine throttling and
        # make completed_pages monotonic for the polling client.
        outcomes = []
        for page in pages:
            ok = one(page.id)
            outcomes.append(ok)
            job = db.get(PptGenerationJob, job_id)
            if job:
                job.completed_pages = sum(1 for item in outcomes if item)
                job.failed_pages = sum(1 for item in outcomes if not item)
                job.current_page_id = page.id
                db.commit()
        job = db.get(PptGenerationJob, job_id)
        job.completed_pages = sum(1 for ok in outcomes if ok); job.failed_pages = sum(1 for ok in outcomes if not ok)
        job.status = "completed" if job.failed_pages == 0 else "completed_with_errors" if job.completed_pages else "failed"; job.finished_at = datetime.now(timezone.utc); db.commit()
    except Exception as exc:
        db.rollback(); job = db.get(PptGenerationJob, job_id)
        if job: job.status, job.error_message = "failed", str(exc)[:500]; job.finished_at = datetime.now(timezone.utc); db.commit()
    finally: db.close()


def _run_generation_job(job_id: str, project_id: str, user_id: str) -> None:
    """Run deck planning and page generation outside the request thread."""
    db = SessionLocal()
    try:
        job = db.get(PptGenerationJob, job_id)
        project = db.get(PptProject, project_id)
        user = db.get(User, user_id)
        if not job or not project or not user or job.status == "cancelled":
            return
        service = PptAgentService(db, user)
        job.status, job.stage, job.started_at = "running", "planning", datetime.now(timezone.utc)
        service._event(project, "generation.started", {"job_id": job.id, "total_pages": job.total_pages}); db.commit()
        pages = sorted(project.pages, key=lambda x: x.sort_order)
        outline = [{"id": p.id, "order": p.sort_order, "section": p.section_title, "title": p.title, "bullets": p.bullets_json or []} for p in pages]
        source_context = service._source_context(project)
        planner_prompt = """你是教学课件 Deck Planner。只输出 JSON：{pages:[{id,role,core_message,content_requirements,visual_type,speaker_notes}]}。大纲只是教学主线，不是页面全文。必须为页面分配 cover、catalogue、section、concept、comparison、process、case、quote、summary、cta 等角色；第一页必须 cover，第二页必须 catalogue。根据主题与页面上下文补充论点、案例、数据和视觉表达建议，不能编造引用。"""
        try:
            planner = ProviderGateway(db, user.id).json(planner_prompt, {"request": project.request_text, "outline": outline, "requirements": source_context})
        except Exception:
            planner = {"pages": []}
        plan_by_id = {str(item.get("id")): item for item in (planner.get("pages") or []) if isinstance(item, dict) and item.get("id")}
        roles = ["cover", "catalogue"]
        for index, page in enumerate(pages):
            item = plan_by_id.get(page.id, {})
            role = str(item.get("role") or (roles[index] if index < 2 else "concept"))
            page.page_role = role
            page.content_plan_json = item
            page.visual_plan_json = {"type": item.get("visual_type") or role}
        job.planner_json = planner if isinstance(planner, dict) else {}; job.stage = "content"; db.commit()
        service._event(project, "planner.completed", {"job_id": job.id}); db.commit()

        layouts = {item["id"]: item for item in __import__("app.services.ppt_theme", fromlist=["list_layouts"]).list_layouts(project.theme_id or "light-academic")}
        layout_list = list(layouts.values())
        assignments = {}
        for index, page in enumerate(pages):
            role = page.page_role
            preferred = "section" if role == "section" else "quote" if role in {"quote", "summary"} else "two-column" if role in {"comparison", "process", "case"} else "title-content"
            assignments[page.id] = next((x["id"] for x in layout_list if x["kind"] in {preferred, "content"} or x["id"] == preferred), layout_list[0]["id"])
        project.layout_assignments = assignments; project.current_stage = "design"; db.commit()

        def render_one(page_id: str):
            local = SessionLocal()
            try:
                p = local.get(PptProject, project_id); u = local.get(User, user_id); page = local.get(PptPage, page_id)
                if not p or not u or not page: return False
                svc = PptAgentService(local, u)
                svc._event(p, "page.content_started", {"job_id": job_id, "page_id": page_id}); local.commit()
                try:
                    svc._visual_research(p, page)
                    svc._visual_plan(p, page)
                    svc._event(p, "page.visual_research_completed", {"job_id": job_id, "page_id": page_id, "asset_count": len((page.visual_plan_json or {}).get("asset_manifest") or [])})
                    local.commit()
                except Exception as exc:
                    # Visual enrichment is best-effort; text-only generation must
                    # remain available when search/image providers are down.
                    page.visual_plan_json = {**(page.visual_plan_json or {}), "asset_status": "degraded", "asset_error": str(exc)[:200]}
                    page.statuses_json = {**(page.statuses_json or {}), "visual": "failed"}
                    local.commit()
                try:
                    combined = ProviderGateway(local, u.id).json("只输出 JSON：{bullets:[string],summary:string,speaker_notes:string,citations:[object],visual_plan:object,elements:[{type,x,y,w,h,text,items,src,font_size,font_weight,color,fill}]}. 根据页面主线与资料补充教学内容并设计版式；不编造引用，坐标使用 0-100 百分比，标题 30-44px、正文 18-24px，避免重叠和文字堆叠。", {"project": p.request_text, "page": svc.serialize_page(page), "plan": page.content_plan_json or {}, "sources": source_context, "layout": layouts.get(assignments.get(page_id), layout_list[0]), "theme": p.theme_config or {}})
                    if isinstance(combined, dict):
                        if isinstance(combined.get("bullets"), list) and combined["bullets"]: page.bullets_json = [str(x) for x in combined["bullets"][:8]]
                        page.summary_md = str(combined.get("summary") or page.summary_md or "")
                        page.speaker_notes = str(combined.get("speaker_notes") or page.speaker_notes or "")
                        if isinstance(combined.get("citations"), list): page.citations_json = combined["citations"][:30]
                        page.visual_plan_json = combined.get("visual_plan") if isinstance(combined.get("visual_plan"), dict) else page.visual_plan_json
                        document = svc._apply_visual_plan(page, svc._normalize_document(combined, layouts.get(assignments.get(page_id), layout_list[0]), p.theme_config or {}))
                    else:
                        raise ValueError("invalid combined design")
                except Exception:
                    try:
                        document = svc._slide_doc(p, page, "design", layout=layouts.get(assignments.get(page_id), layout_list[0]), outline=outline)
                    except Exception:
                        document = svc._template_document(page, p.theme_config or {}, layouts.get(assignments.get(page_id), layout_list[0]))
                svc._event(p, "page.content_completed", {"job_id": job_id, "page_id": page_id}); local.commit()
                layout = layouts.get(assignments.get(page_id), layout_list[0])
                if not svc._document_quality_ok(document):
                    try:
                        repaired = svc._slide_doc(p, page, "design_repair", layout=layout, outline=outline)
                        document = repaired if svc._document_quality_ok(repaired) else svc._template_document(page, p.theme_config or {}, layout)
                    except Exception:
                        document = svc._template_document(page, p.theme_config or {}, layout)
                page.design_document_json = document; page.document_revision += 1; page.statuses_json = {**(page.statuses_json or {}), "design": "ready"}
                version_no = (local.scalar(select(PptDocumentVersion).where(PptDocumentVersion.page_id == page.id).order_by(PptDocumentVersion.version_no.desc())) or PptDocumentVersion(version_no=0)).version_no + 1
                version = PptDocumentVersion(project_id=p.id, page_id=page.id, version_no=version_no, document_json=document)
                local.add(version); local.flush(); page.current_document_version_id = version.id
                local.commit()
                progress = local.scalar(select(PptGenerationJob).where(PptGenerationJob.id == job_id).with_for_update())
                if progress:
                    progress.completed_pages = min(progress.total_pages, (progress.completed_pages or 0) + 1)
                    progress.current_page_id = page_id
                    completed = progress.completed_pages
                    total = progress.total_pages
                else:
                    completed = 0
                    total = len(p.pages)
                svc._event(p, "page.designed", {"job_id": job_id, "page_id": page_id, "layout_id": layout["id"], "completed_pages": completed, "total_pages": total}); local.commit()
                return True
            finally:
                local.close()

        workers = max(1, min(int(getattr(settings, "ppt_generation_workers", 4)), 8))
        with ThreadPoolExecutor(max_workers=workers) as pool:
            results = list(pool.map(render_one, [p.id for p in pages]))
        job = db.get(PptGenerationJob, job_id)
        job.completed_pages = sum(1 for ok in results if ok); job.failed_pages = len(results) - job.completed_pages; job.stage = "quality"; db.commit()
        service._event(project, "generation.quality_checked", {"job_id": job.id, "completed_pages": job.completed_pages, "failed_pages": job.failed_pages});
        job.status = "completed" if not job.failed_pages else "completed_with_errors"; job.stage = "completed"; job.finished_at = datetime.now(timezone.utc); project.design_status = "ready"; project.current_stage = "design"; db.commit()
        service._event(project, "generation.completed", {"job_id": job.id, "completed_pages": job.completed_pages, "failed_pages": job.failed_pages}); db.commit()
    except Exception as exc:
        db.rollback(); job = db.get(PptGenerationJob, job_id)
        if job:
            job.status, job.stage, job.error_message = "failed", "failed", str(exc); job.finished_at = datetime.now(timezone.utc); db.commit()
            project = db.get(PptProject, project_id)
            if project:
                PptAgentService(db, db.get(User, user_id))._event(project, "generation.failed", {"job_id": job_id, "error": str(exc)}); db.commit()
    finally:
        db.close()




