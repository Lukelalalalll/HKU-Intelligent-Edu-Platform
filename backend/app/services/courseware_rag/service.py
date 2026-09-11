from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

from sqlalchemy import select, delete
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.session import SessionLocal
from app.models import Course, CourseChapter, CourseMaterial, CourseMaterialChunk, CourseMaterialIngestion, Enrollment
from app.services.ai_gateway import gateway_for

_executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="courseware-rag")

def _read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        except Exception:
            return ""
    if suffix == ".docx":
        try:
            from docx import Document
            return "\n".join(p.text for p in Document(str(path)).paragraphs)
        except Exception:
            return ""
    if suffix == ".pptx":
        try:
            from pptx import Presentation
            lines = []
            for slide_no, slide in enumerate(Presentation(str(path)).slides, 1):
                text = " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
                if text.strip(): lines.append(f"[Page {slide_no}] {text}")
            return "\n".join(lines)
        except Exception:
            return ""
    if suffix in {".xlsx", ".xls"}:
        try:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            return "\n".join(" | ".join(str(v or "") for v in row) for ws in book.worksheets for row in ws.iter_rows(values_only=True))
        except Exception:
            return ""
    return ""

def _chunks(text: str, size: int = 1200, overlap: int = 160):
    text = re.sub(r"\s+", " ", text).strip()
    if not text: return []
    step = max(1, size - overlap)
    return [text[i:i + size] for i in range(0, len(text), step)]

def ingest_material(material_id: str) -> None:
    db = SessionLocal()
    try:
        material = db.scalar(select(CourseMaterial).where(CourseMaterial.id == material_id))
        if not material: return
        ingestion = db.scalar(select(CourseMaterialIngestion).where(CourseMaterialIngestion.material_id == material_id))
        if not ingestion:
            ingestion = CourseMaterialIngestion(material_id=material_id)
            db.add(ingestion)
        ingestion.status = "processing"; ingestion.started_at = datetime.now(timezone.utc); ingestion.error_message = None
        db.commit()
        path = settings.upload_path / material.file_asset.storage_key
        text = _read_document(path)
        db.execute(delete(CourseMaterialChunk).where(CourseMaterialChunk.material_id == material_id))
        for index, content in enumerate(_chunks(text)):
            page = None
            match = re.search(r"\[Page (\d+)\]", content)
            if match: page = int(match.group(1))
            db.add(CourseMaterialChunk(course_id=material.chapter.course_id, chapter_id=material.chapter_id, material_id=material.id, chunk_id=f"{material.id}:{index}", content=content, page_number=page, source_anchor=f"chunk-{index}"))
        ingestion.status = "ready"; ingestion.completed_at = datetime.now(timezone.utc)
        db.commit()
    except Exception as exc:
        db.rollback()
        if 'ingestion' in locals() and ingestion:
            ingestion.status = "failed"; ingestion.error_message = str(exc)[:1000]; db.commit()
    finally:
        db.close()

def queue_ingestion(material_id: str):
    _executor.submit(ingest_material, material_id)

def query_courses(db: Session, question: str, student_id: str, course_id: str | None = None):
    courses = db.scalars(select(Course).join(Enrollment).where(Enrollment.student_id == student_id)).all()
    if course_id:
        courses = [c for c in courses if c.id == course_id]
    terms = set(re.findall(r"[\w-]+", question.casefold()))
    scored = []
    for course in courses:
        score = sum(1 for t in terms if t in f"{course.code} {course.name}".casefold())
        scored.append((score, course))
    scored.sort(key=lambda x: x[0], reverse=True)
    if not scored: return None, [], []
    if not course_id and scored[0][0] == 0 and len(scored) > 1:
        return None, scored[:3], []
    selected = scored[0][1]
    chunks = db.scalars(select(CourseMaterialChunk).where(CourseMaterialChunk.course_id == selected.id)).all()
    ranked = sorted(chunks, key=lambda chunk: sum(1 for term in terms if term in chunk.content.casefold()), reverse=True)[:5]
    return selected, [], ranked

def courseware_rag(db: Session, question: str, student_id: str, course_id: str | None = None):
    course, candidates, chunks = query_courses(db, question, student_id, course_id)
    if not course:
        return {"course_id": None, "course": None, "answer": "请先选择一门课程，我会基于该课程的课件为你讲解。", "citations": [], "confidence": 0.0, "needs_course_selection": bool(candidates), "candidate_courses": [{"id": c.id, "code": c.code, "name": c.name} for _, c in candidates]}
    if not chunks:
        answer = "这门课程目前还没有可用的已索引课件。请稍后再试，或联系教师完成材料索引。"
    else:
        context = "\n\n".join(chunk.content for chunk in chunks)
        answer = gateway_for(db, "student_lecturer").chat(
            "你是课程 AI 讲师。只能依据提供的课程材料回答，清楚解释概念并在不确定时说明。",
            f"课程：{course.code}《{course.name}》\n课程材料：\n{context[:10000]}\n\n学生问题：{question}",
        )
    citations = []
    for chunk in chunks:
        material = db.get(CourseMaterial, chunk.material_id); chapter = db.get(CourseChapter, chunk.chapter_id)
        if material and chapter:
            citations.append({"course_id": course.id, "course_code": course.code, "course_name": course.name, "chapter_id": chapter.id, "chapter_title": chapter.title, "material_id": material.id, "material_title": material.title, "page_number": chunk.page_number, "href": f"/courses/{course.id}?chapter={chapter.id}&material={material.id}"})
    gateway = gateway_for(db, "student_lecturer") if chunks else None
    return {"course_id": course.id, "course": {"id": course.id, "code": course.code, "name": course.name}, "answer": answer, "citations": citations, "confidence": 0.85 if chunks else 0.1, "needs_course_selection": False, "candidate_courses": [], "model": gateway.model if gateway else None}
