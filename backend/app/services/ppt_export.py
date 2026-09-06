from __future__ import annotations

import base64
import re
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote_to_bytes

from fastapi import HTTPException


SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5


def _rgb(value: str | None, fallback: tuple[int, int, int] = (15, 45, 35)) -> tuple[int, int, int]:
    """Parse #rgb/#rrggbb values without letting malformed editor data abort export."""
    if not value:
        return fallback
    raw_value = str(value).strip()
    # Accept the CSS colour forms emitted by the browser/editor as well as
    # hexadecimal values. Alpha is intentionally ignored because python-pptx
    # stores transparency separately from RGB.
    css = re.fullmatch(r"rgba?\(\s*([\d.]+)[,\s]+([\d.]+)[,\s]+([\d.]+)(?:[,\s/]+[\d.]+)?\s*\)", raw_value, re.I)
    if css:
        try:
            return tuple(max(0, min(255, int(float(css.group(index))))) for index in (1, 2, 3))
        except (TypeError, ValueError):
            return fallback
    raw = raw_value.lstrip("#")
    if len(raw) == 3:
        raw = "".join(ch * 2 for ch in raw)
    if len(raw) != 6:
        return fallback
    try:
        return tuple(int(raw[index:index + 2], 16) for index in (0, 2, 4))
    except ValueError:
        return fallback


def _number(value: Any, fallback: float = 0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return fallback


def _length(percent: Any, total: int) -> int:
    return int(max(0.0, min(100.0, _number(percent))) / 100 * total)


def _geometry(element: dict[str, Any], prs) -> tuple[int, int, int, int]:
    left = _length(element.get("x"), prs.slide_width)
    top = _length(element.get("y"), prs.slide_height)
    width = max(1, _length(element.get("w", 20), prs.slide_width))
    height = max(1, _length(element.get("h", 10), prs.slide_height))
    return left, top, width, height


def _element_text(element: dict[str, Any]) -> str:
    text = element.get("text")
    if text is not None:
        return str(text)
    return "\n".join(str(item) for item in element.get("items") or [])


def _first_value(element: dict[str, Any], *keys: str) -> Any:
    """Return the first non-empty value, supporting editor field aliases."""
    for key in keys:
        value = element.get(key)
        if value is not None and str(value).strip() != "":
            return value
    return None


VISUAL_TYPES = {
    "shape", "rect", "rectangle", "rounded-rectangle", "rounded_rectangle",
    "roundrect", "card", "panel", "box", "bar", "divider", "line",
    "rule", "circle", "oval", "ellipse", "background", "decoration",
}


def _is_visual_element(element: dict[str, Any]) -> bool:
    element_type = str(element.get("type") or "").strip().lower().replace(" ", "-")
    if element_type in VISUAL_TYPES:
        return True
    # Some model responses omit the type for decorative geometry. Do not
    # interpret ordinary text boxes as shapes unless they have no text and
    # carry an explicit paint/shape hint.
    if not _element_text(element) and any(_first_value(element, key) is not None for key in ("fill", "fill_color", "background", "background_color", "stroke", "stroke_color", "shape")):
        return True
    return False


def _font_is_bold(value: Any, element_type: str) -> bool:
    if value is None:
        return element_type == "title"
    if isinstance(value, (int, float)):
        return value >= 600
    return str(value).strip().lower() in {"bold", "bolder", "600", "700", "800", "900"}


def _alignment(value: Any, PP_ALIGN):
    return {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(str(value or "left").lower(), PP_ALIGN.LEFT)


def _data_image(source: str) -> BytesIO | Path | None:
    """Return a python-pptx compatible image source without fetching remote URLs."""
    if source.startswith("data:"):
        try:
            header, payload = source.split(",", 1)
            content = base64.b64decode(payload) if ";base64" in header.lower() else unquote_to_bytes(payload)
            return BytesIO(content)
        except (ValueError, UnicodeError):
            return None
    path = Path(source)
    return path if path.is_file() else None


def _fallback_elements(page, colors: dict[str, str]) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = [{
        "id": "title", "type": "title", "x": 7, "y": 8, "w": 86, "h": 14,
        "text": page.title or "未命名页面", "font_size": 30, "color": colors.get("title"),
    }]
    for index, bullet in enumerate((page.bullets_json or [])[:6]):
        elements.append({
            "id": f"bullet-{index}", "type": "body", "x": 9, "y": 30 + index * 10, "w": 80, "h": 8,
            "text": f"• {bullet}", "font_size": 18, "color": colors.get("body"),
        })
    return elements


def _page_document(page) -> dict[str, Any]:
    document = getattr(page, "design_document_json", None) or getattr(page, "draft_document_json", None) or {}
    return document if isinstance(document, dict) else {}


def _page_colors(theme: dict[str, Any], document: dict[str, Any]) -> dict[str, str]:
    colors = dict(theme.get("colors") or {})
    document_theme = document.get("theme")
    if isinstance(document_theme, dict):
        colors.update({key: str(value) for key, value in document_theme.items() if key in {"bg", "surface", "title", "body", "accent"} and value})
        for alias in ("background", "background_color"):
            if document_theme.get(alias):
                colors["bg"] = str(document_theme[alias])
        if document_theme.get("surface_color"):
            colors["surface"] = str(document_theme["surface_color"])
    return colors


def _add_text(slide, element: dict[str, Any], colors: dict[str, str], prs, RGBColor, Pt, PP_ALIGN) -> None:
    text = _element_text(element)
    if not text:
        return
    left, top, width, height = _geometry(element, prs)
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.fill.background()
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    element_type = str(element.get("type") or "body")
    font_size = max(6, _number(element.get("font_size"), 30 if element_type == "title" else 18))
    color = RGBColor(*_rgb(element.get("color") or colors.get("title" if element_type == "title" else "body")))
    for index, line in enumerate(text.splitlines() or [""]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = _alignment(element.get("align"), PP_ALIGN)
        for run in paragraph.runs:
            run.font.name = str(element.get("font_family") or "Arial")
            run.font.size = Pt(font_size)
            run.font.bold = _font_is_bold(element.get("font_weight"), element_type)
            run.font.color.rgb = color


def _add_shape(slide, element: dict[str, Any], colors: dict[str, str], prs, RGBColor, MSO_SHAPE, Pt) -> None:
    left, top, width, height = _geometry(element, prs)
    element_type = str(element.get("type") or "shape").strip().lower().replace(" ", "-")
    shape_hint = str(element.get("shape") or "").strip().lower().replace("_", "-")
    if element_type in {"circle", "oval", "ellipse"} or shape_hint in {"circle", "oval", "ellipse"}:
        shape_type = MSO_SHAPE.OVAL
    elif element_type in {"rounded-rectangle", "rounded_rectangle", "roundrect", "card", "panel"} or shape_hint in {"rounded-rectangle", "roundrect", "card", "panel"} or _number(element.get("radius")) > 0:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        shape_type = MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    fill_value = _first_value(element, "fill", "fill_color", "background", "background_color")
    stroke_value = _first_value(element, "stroke", "stroke_color", "border_color", "line_color")
    # Thin rules/lines are represented as filled rectangles so they remain
    # reliable across PowerPoint versions and preserve rounded end caps when
    # the source used a radius.
    if (element_type in {"line", "divider", "rule", "bar"} or shape_hint in {"line", "divider", "rule", "bar"}) and fill_value is None:
        fill_value = stroke_value or colors.get("accent")
    if fill_value is None:
        fill_value = colors.get("surface") or colors.get("accent")
    if fill_value is None or str(fill_value).lower() in {"none", "transparent"}:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*_rgb(str(fill_value), _rgb(colors.get("surface"))))
    if stroke_value is None or str(stroke_value).lower() in {"none", "transparent"}:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor(*_rgb(str(stroke_value), _rgb(colors.get("accent"))))
        shape.line.width = Pt(max(0.5, _number(_first_value(element, "stroke_width", "line_width"), 1)))
    opacity = _number(element.get("opacity"), 1.0)
    if 0 <= opacity < 1:
        try:
            shape.fill.transparency = int((1 - opacity) * 100)
        except (AttributeError, ValueError):
            pass


def _add_image(slide, element: dict[str, Any], prs) -> None:
    source = _data_image(str(element.get("asset_path") or element.get("src") or ""))
    if source is None:
        return
    left, top, width, height = _geometry(element, prs)
    try:
        slide.shapes.add_picture(source, left, top, width=width, height=height)
    except Exception:
        # A malformed or unsupported uploaded image must not prevent the rest
        # of a teacher's deck from exporting.
        return


def export_project_pptx(project, output: Path, filename: str | None = None) -> None:
    """Export the saved page documents as native PowerPoint objects.

    Browser previews use SVG, but python-pptx/Pillow cannot embed SVG streams.
    Creating native text, shape, and image objects avoids platform-specific SVG
    rasterizer dependencies (notably Cairo DLLs on Windows) and keeps exports
    editable in PowerPoint.
    """
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        from app.services.ppt_theme import get_theme
    except ImportError as exc:
        raise HTTPException(500, "导出 PPTX 需要 python-pptx，请重新安装 requirements.txt") from exc

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_WIDTH_INCHES), Inches(SLIDE_HEIGHT_INCHES)
    blank = prs.slide_layouts[6]
    theme = get_theme(getattr(project, "theme_id", None) or "light-academic")

    for page in sorted(project.pages, key=lambda item: item.sort_order):
        slide = prs.slides.add_slide(blank)
        document = _page_document(page)
        colors = _page_colors(theme, document)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(*_rgb(colors.get("bg"), (251, 254, 252)))
        elements = document.get("elements") if isinstance(document.get("elements"), list) else None
        elements = elements or _fallback_elements(page, colors)

        for element in sorted((item for item in elements if isinstance(item, dict)), key=lambda item: _number(item.get("zIndex"), 1)):
            element_type = str(element.get("type") or "body")
            if element_type == "image" or (element_type == "icon" and element.get("src")):
                _add_image(slide, element, prs)
            elif _is_visual_element(element):
                _add_shape(slide, element, colors, prs, RGBColor, MSO_SHAPE, Pt)
            else:
                _add_text(slide, element, colors, prs, RGBColor, Pt, PP_ALIGN)

        if page.speaker_notes:
            try:
                slide.notes_slide.notes_text_frame.text = page.speaker_notes
            except AttributeError:
                # Older python-pptx versions do not expose notes editing.
                pass

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
