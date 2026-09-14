from .context import *  # noqa: F401,F403

class _DocumentMixin:
    def _remote_image_asset(url: str, page: PptPage, index: int) -> dict[str, Any] | None:
        """Return a browser-renderable fallback after local caching fails.

        Only public HTTP(S) hosts are allowed here; this mirrors the SSRF
        guard used by ``_download_image`` and prevents a failed cache request
        from turning into an unsafe remote URL in the client.
        """
        try:
            validate_public_url(url)
        except ValueError:
            return None
        return {
            "id": f"asset-{page.id}-{index}",
            "src": url,
            "public_url": url,
            "source_url": url,
            "status": "ready",
            "remote_only": True,
        }


    def _visual_plan(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        visual = dict(page.visual_plan_json or {})
        all_assets = [item for item in visual.get("asset_manifest", []) if isinstance(item, dict) and item.get("status") == "ready"]
        selected = visual.get("selected_asset_ids")
        assets = [item for item in all_assets if not (isinstance(selected, list) and item.get("id") not in selected)]
        role = page.page_role or "concept"
        fallback_kind = {"cover": "image-focus", "section": "image-focus", "case": "image-text", "process": "process-timeline", "comparison": "two-column"}.get(role, "image-text" if assets else "content")
        plan = {"layout_kind": fallback_kind, "image_slots": []}
        for index, asset in enumerate(assets[:6]):
            if index == 0 and role in {"cover", "section"}:
                x, y, w, h, hint = 0, 0, 100, 100, "background"
            elif len(assets) == 1:
                x, y, w, h, hint = 54, 18, 40, 70, "right"
            elif len(assets) <= 3:
                x, y, w, h, hint = 54 + (index % 2) * 21, 22 + (index // 2) * 34, 19, 28, "supporting"
            else:
                x, y, w, h, hint = 54 + (index % 3) * 15, 18 + (index // 3) * 34, 13, 28, "grid"
            plan["image_slots"].append({"slot_id": f"image-{index + 1}", "asset_id": asset["id"], "src": asset.get("public_url") or asset.get("src"), "asset_path": asset.get("asset_path"), "placement_hint": hint, "x": x, "y": y, "w": w, "h": h, "object_fit": "cover", "confidence": round(min(0.98, 0.62 + float(asset.get("score") or 0) * 0.3), 2)})
        visual.update(plan)
        page.visual_plan_json = visual
        page.statuses_json = {**(page.statuses_json or {}), "visual": "ready"}
        self.db.commit()
        return visual


    def _summary(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        text = "\n\n".join(c.content_md for col in self.db.scalars(select(PptSourceCollection).where(PptSourceCollection.project_id == p.id, PptSourceCollection.page_id == page.id)) for d in col.documents for c in d.chunks)
        if not text and not page.citations_json: raise HTTPException(409, "当前页没有资料，无法生成 Summary")
        gateway = ProviderGateway(self.db, self.user.id); page.summary_md = gateway.text("根据给定资料生成有引用意识的教学 Summary，使用 Markdown，禁止编造事实。", {"title": page.title, "bullets": page.bullets_json, "sources": text or page.citations_json}); page.statuses_json = {**(page.statuses_json or {}), "summary": "ready"}; self.db.commit(); return {"summary_length": len(page.summary_md)}


    def _template_document(self, page: PptPage, theme: dict[str, Any], layout: dict[str, Any]) -> dict[str, Any]:
        colors = theme.get("colors") or {}
        bullets = [str(item) for item in (page.bullets_json or [])[:6]]
        kind = layout.get("kind", "content")
        role = page.page_role or "concept"
        title_size = 42 if role in {"cover", "section"} else 32
        elements: list[dict[str, Any]] = [{"id": "title", "type": "title", "x": 7, "y": 8, "w": 86, "h": 14, "text": page.title, "font_size": title_size, "font_weight": 800, "color": colors.get("title")}]
        if role == "cover":
            elements.append({"id": "cover-subtitle", "type": "caption", "x": 9, "y": 38, "w": 76, "h": 12, "text": page.section_title or "课程课件", "font_size": 22, "font_weight": 500, "color": colors.get("body")})
        elif role == "catalogue":
            for index, text in enumerate(bullets[:8]):
                elements.append({"id": f"catalogue-{index}", "type": "body", "x": 10, "y": 29 + index * 7, "w": 78, "h": 6, "text": f"{index + 1:02d}  {text}", "font_size": 20, "font_weight": 600 if index < 3 else 500, "color": colors.get("body")})
        elif role == "section" or kind == "section":
            elements.append({"id": "section", "type": "body", "x": 10, "y": 43, "w": 80, "h": 16, "text": page.section_title or page.title, "font_size": 34, "font_weight": 700, "color": colors.get("title")})
        elif kind == "quote":
            elements.append({"id": "quote", "type": "body", "x": 12, "y": 38, "w": 76, "h": 24, "text": bullets[0] if bullets else page.title, "font_size": 26, "color": colors.get("body")})
        elif kind == "columns":
            midpoint = max(1, (len(bullets) + 1) // 2)
            for index, text in enumerate(bullets):
                column, row = (0, index) if index < midpoint else (1, index - midpoint)
                elements.append({"id": f"bullet-{index}", "type": "body", "x": 9 + column * 44, "y": 29 + row * 12, "w": 38, "h": 9, "text": f"• {text}", "font_size": 18, "font_weight": 500, "color": colors.get("body")})
        else:
            for index, text in enumerate(bullets):
                elements.append({"id": f"bullet-{index}", "type": "body", "x": 9, "y": 29 + index * 10, "w": 82, "h": 8, "text": f"• {text}", "font_size": 18 if len(text) < 70 else 15, "font_weight": 500, "color": colors.get("body")})
        visual = page.visual_plan_json or {}
        for index, slot in enumerate(visual.get("image_slots") or []):
            if slot.get("src"):
                elements.append({"id": f"image-slot-{index + 1}", "type": "image", "asset_id": slot.get("asset_id"), "src": slot.get("src"), "asset_path": slot.get("asset_path"), "alt": page.title, "x": slot.get("x", 54), "y": slot.get("y", 18), "w": slot.get("w", 40), "h": slot.get("h", 70), "object_fit": slot.get("object_fit", "cover"), "zIndex": 0})
        editor_theme = {**colors, "background": colors.get("bg")}
        return {"version": 2, "canvas": {"width": 1280, "height": 720}, "layout": layout["id"], "theme": editor_theme, "elements": elements, "speaker_notes": page.speaker_notes}


    @staticmethod
    def _document_quality_ok(document: dict[str, Any]) -> bool:
        elements = document.get("elements") if isinstance(document, dict) else None
        if not isinstance(elements, list) or not elements:
            return False
        title_sizes = [float(item.get("font_size", 0)) for item in elements if isinstance(item, dict) and str(item.get("type", "")).lower() in {"title", "heading"}]
        max_title = max(title_sizes, default=0)
        for element in elements:
            if not isinstance(element, dict) or not str(element.get("text") or element.get("src") or "").strip():
                continue
            try:
                x, y, w, h = (float(element.get(key, 0)) for key in ("x", "y", "w", "h"))
            except (TypeError, ValueError):
                return False
            if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > 100.5 or y + h > 100.5:
                return False
            text = str(element.get("text") or "")
            size = float(element.get("font_size", 18) or 18)
            if text and len(text) > 220 and h < 12:
                return False
            if str(element.get("type", "")).lower() in {"body", "caption"} and max_title and size > max_title:
                return False
        return True


    def _slide_doc(self, p: PptProject, page: PptPage, mode: str, layout: dict[str, Any] | None = None, outline: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        gateway = ProviderGateway(self.db, self.user.id)
        selected_layout = layout or {"id": (page.document or {}).get("layout", "title-content"), "kind": "content"}
        raw = gateway.json("你是资深教学课件设计 Agent。只输出结构化 Slide JSON：{layout,theme,elements:[{type,x,y,w,h,text,items,src,font_size,font_weight,color,fill}],speaker_notes}。大纲只是页面职责，不是最终全文；必须优先使用 content_plan 扩写后的论点、案例和视觉建议。根据页面角色设计清晰层级：标题 30-44px、正文 18-24px、辅助文字 12-16px；封面/目录/章节/对比/流程/案例/总结使用匹配的视觉结构。所有坐标使用 0-100 百分比，留出安全边距，禁止严重重叠和文字堆叠。", {"mode": mode, "page": {"title": page.title, "section": page.section_title, "role": page.page_role, "bullets": page.bullets_json, "summary": page.summary_md, "content_plan": page.content_plan_json or {}, "visual_plan": page.visual_plan_json or {}}, "outline": outline or [], "layout": selected_layout, "theme": p.theme_config or {}})
        return self._apply_visual_plan(page, self._normalize_document(raw, selected_layout, p.theme_config or {}))


    def _apply_visual_plan(self, page: PptPage, document: dict[str, Any]) -> dict[str, Any]:
        result = dict(document or {})
        elements = [dict(item) for item in result.get("elements") or [] if isinstance(item, dict)]
        existing_asset_ids = {str(item.get("asset_id")) for item in elements if item.get("asset_id")}
        plan = page.visual_plan_json or {}
        for index, slot in enumerate(plan.get("image_slots") or []):
            if not isinstance(slot, dict) or not slot.get("src") or str(slot.get("asset_id")) in existing_asset_ids:
                continue
            elements.append({"id": f"image-slot-{index + 1}", "type": "image", "asset_id": slot.get("asset_id"), "src": slot.get("src"), "asset_path": slot.get("asset_path"), "alt": slot.get("alt") or page.title, "x": slot.get("x", 54), "y": slot.get("y", 18), "w": slot.get("w", 40), "h": slot.get("h", 70), "object_fit": slot.get("object_fit", "cover"), "zIndex": 0})
        result["elements"] = elements
        result["visual_plan"] = plan
        return self._repair_document_layout(result)


    @staticmethod
    def _repair_document_layout(document: dict[str, Any]) -> dict[str, Any]:
        """Apply conservative geometry repairs after model/layout composition."""
        elements = [dict(item) for item in document.get("elements") or [] if isinstance(item, dict)]
        images = [item for item in elements if str(item.get("type") or "").lower() == "image" and float(item.get("x", 0) or 0) > 0]
        text_items = [item for item in elements if str(item.get("type") or "").lower() in {"body", "caption"}]
        for image in images:
            ix, iy, iw, ih = (float(image.get(key, 0) or 0) for key in ("x", "y", "w", "h"))
            for text in text_items:
                tx, ty, tw, th = (float(text.get(key, 0) or 0) for key in ("x", "y", "w", "h"))
                intersects = ix < tx + tw and ix + iw > tx and iy < ty + th and iy + ih > ty
                if not intersects:
                    continue
                if ix >= 50:
                    text["w"] = max(8, min(tw, ix - tx - 3))
                else:
                    image["x"] = min(96 - iw, max(ix, tx + tw + 3))
        for item in elements:
            for key in ("x", "y", "w", "h"):
                try:
                    item[key] = max(0.0, min(100.0, float(item.get(key, 0))))
                except (TypeError, ValueError):
                    item[key] = 0.0
            item["w"] = min(item["w"], 100 - item["x"]); item["h"] = min(item["h"], 100 - item["y"])
        return {**document, "elements": elements}


    @staticmethod
    def _normalize_document(raw: dict[str, Any], selected_layout: dict[str, Any], theme_config: dict[str, Any]) -> dict[str, Any]:
        normalized = dict(raw or {}); normalized["version"] = 2; normalized["canvas"] = {"width": 1280, "height": 720}; normalized["elements"] = []
        normalized["layout"] = selected_layout["id"]; colors = (theme_config or {}).get("colors", {}); normalized["theme"] = {**colors, "background": colors.get("bg"), "tokens": (theme_config or {}).get("tokens", {}), "accessibility": (theme_config or {}).get("accessibility", {})}
        for index, item in enumerate(raw.get("elements") or []):
            element = dict(item or {}); element["id"] = str(element.get("id") or f"element-{index + 1}")
            for key in ("x", "y", "w", "h"):
                try: element[key] = max(0.0, min(100.0, float(element.get(key, 0))))
                except (TypeError, ValueError): element[key] = 0.0
            element.setdefault("type", "body")
            element.setdefault("zIndex", index + 1)
            if element.get("type") == "image": element.setdefault("object_fit", "cover")
            normalized["elements"].append(element)
        return normalized


    def _draft(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        if not page.summary_md: raise HTTPException(409, "请先生成当前页 Summary")
        page.draft_document_json = self._slide_doc(p, page, "draft"); page.statuses_json = {**(page.statuses_json or {}), "draft": "ready"}; self.db.commit(); return {"page_id": page.id, "document": page.draft_document_json}


    def _design(self, p: PptProject, page: PptPage) -> dict[str, Any]:
        if not page.draft_document_json: raise HTTPException(409, "请先生成当前页 Draft")
        page.design_document_json = self._slide_doc(p, page, "design"); page.document_revision += 1
        version_no = (self.db.scalar(select(PptDocumentVersion).where(PptDocumentVersion.page_id == page.id).order_by(PptDocumentVersion.version_no.desc())) or PptDocumentVersion(version_no=0)).version_no + 1
        version = PptDocumentVersion(project_id=p.id, page_id=page.id, version_no=version_no, document_json=page.design_document_json)
        self.db.add(version); self.db.flush(); page.current_document_version_id = version.id
        page.statuses_json = {**(page.statuses_json or {}), "design": "ready"}; self.db.commit(); return {"page_id": page.id, "document": page.design_document_json, "revision": page.document_revision}


    def _batch(self, p: PptProject, action: str) -> dict[str, Any]:
        mapping = {"project_batch_search": "page_search_run", "project_batch_visual": "page_visual_research", "project_batch_summary": "page_summary_generate", "project_batch_draft": "page_draft_generate", "project_batch_design": "page_design_generate"}; results = []
        for page in sorted(p.pages, key=lambda x: x.sort_order):
            try: results.append(self.run_action(p.id, page.id, mapping[action])["result"])
            except HTTPException as exc: results.append({"page_id": page.id, "status": "skipped", "reason": exc.detail})
        return {"items": results}


    def upload_document(self, project_id: str, page_id: str | None, upload: UploadFile) -> dict[str, Any]:
        p = self.project(project_id); page = self.page(project_id, page_id) if page_id else None
        content = upload.file.read()
        try:
            asset, processed, job = upload_asset(self.db, uploader_id=self.user.id, filename=upload.filename or "document", mime_type=upload.content_type, content=content, target_type="ppt_project", target_id=p.id, owner_id=self.user.id, visibility="owner", metadata={"page_id": page.id if page else None})
        except ValueError as exc: raise HTTPException(415, str(exc)) from exc
        name = asset.original_name; root = settings.ppt_storage_path / p.id; root.mkdir(parents=True, exist_ok=True); path = root / name; path.write_bytes(content)
        text = self._extract_text(path, content); collection = self.db.scalar(select(PptSourceCollection).where(PptSourceCollection.project_id == p.id, PptSourceCollection.page_id == (page.id if page else None)))
        if not collection: collection = PptSourceCollection(project_id=p.id, page_id=page.id if page else None, collection_type="page_corpus" if page else "init_corpus", title="页面资料池" if page else "项目资料池"); self.db.add(collection); self.db.flush()
        doc = PptSourceDocument(collection_id=collection.id, source_type="upload", source_uri=str(path), title=name, content_md=text, metadata_json={"mime_type": upload.content_type}); self.db.add(doc); self.db.flush()
        chunks = self._chunk_text(text); doc.chunks = [PptSourceChunk(chunk_index=i, content_md=chunk) for i, chunk in enumerate(chunks)]; self.db.commit(); bind_document(self.db, processed.id, target_type="ppt_project", target_id=p.id, owner_id=self.user.id, visibility="owner", metadata={"page_id": page.id if page else None}); return {"document_id": doc.id, "processing_document_id": processed.id, "job_id": job.id if job else None, "title": name, "chunk_count": len(chunks), "collection_type": collection.collection_type}


    def _chunk_text(self, text: str) -> list[str]:
        try:
            from langchain_text_splitters import RecursiveCharacterTextSplitter
            return RecursiveCharacterTextSplitter(chunk_size=1800, chunk_overlap=200).split_text(text) or [""]
        except ImportError:
            return [text[i:i + 1800] for i in range(0, len(text), 1600)] or [""]


    def _extract_text(self, path: Path, content: bytes) -> str:
        if path.suffix.lower() in {".txt", ".md", ".markdown", ".csv"}: return content.decode("utf-8", errors="ignore")
        if path.suffix.lower() == ".json":
            try: return json.dumps(json.loads(content.decode("utf-8", errors="ignore")), ensure_ascii=False, indent=2)
            except Exception: return content.decode("utf-8", errors="ignore")
        if path.suffix.lower() in {".xlsx", ".xls"}:
            try:
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=True)
                return "\n".join(" | ".join(str(value or "") for value in row) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True))
            except Exception: return ""
        if path.suffix.lower() == ".pdf":
            try:
                import pypdf
                return "\n".join((page.extract_text() or "") for page in pypdf.PdfReader(str(path)).pages)
            except Exception: return "PDF 已上传，统一文件处理任务正在生成页级内容。"
        if path.suffix.lower() == ".docx":
            try:
                from docx import Document
                return "\n".join(p.text for p in Document(str(path)).paragraphs)
            except Exception: return "DOCX 已上传，统一文件处理任务正在生成结构化内容。"
        if path.suffix.lower() == ".pptx":
            try:
                from pptx import Presentation
                return "\n".join(shape.text for slide in Presentation(str(path)).slides for shape in slide.shapes if hasattr(shape, "text"))
            except Exception: return "PPTX 已上传，统一文件处理任务正在生成幻灯片内容。"
        if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
            # Vision is best-effort: the original image remains available even
            # when the configured provider does not support image inputs.
            try:
                import base64
                gateway = ProviderGateway(self.db, self.user.id)
                client = gateway._client()
                model = gateway.config.model if gateway.config else settings.llm_model
                mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}[path.suffix.lower()]
                response = client.chat.completions.create(model=model, temperature=0.1, messages=[{"role": "user", "content": [{"type": "text", "text": "请提取图片中的文字，并描述图表、关键视觉信息和可用于教学课件的要点。"}, {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{base64.b64encode(content).decode()}"}}]}])
                return response.choices[0].message.content or "图片资料（未提取到文字）"
            except Exception:
                return "图片资料已上传；当前模型不支持视觉解析，请教师在对话中补充图片要点。"
        return ""

