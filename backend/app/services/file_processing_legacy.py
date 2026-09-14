from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.session import SessionLocal
from app.models import Assignment, CourseMaterial, CourseMaterialIngestion, FileAsset, FileContextBinding, FileProcessingChunk, FileProcessingDocument, FileProcessingJob, FileProcessingPage, Submission

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".xlsx", ".xls", ".txt", ".md", ".markdown", ".csv", ".json", ".png", ".jpg", ".jpeg", ".webp"}


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def _safe_name(filename: str | None) -> str:
    name = Path(filename or "upload").name.strip() or "upload"
    return name[:500]


def validate_upload(filename: str | None, content_type: str | None, size: int) -> tuple[str, str]:
    name = _safe_name(filename)
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("仅支持 PDF、PPTX、DOCX、XLSX、TXT、MD、CSV、JSON 和图片文件")
    if size <= 0:
        raise ValueError("文件不能为空")
    if size > settings.file_processing_max_bytes:
        raise ValueError(f"文件不能超过 {settings.file_processing_max_bytes // 1024 // 1024}MB")
    expected = {".pdf": {"application/pdf"}, ".json": {"application/json", "text/json", "text/plain"}, ".csv": {"text/csv", "text/plain"}, ".txt": {"text/plain"}, ".md": {"text/markdown", "text/plain"}, ".markdown": {"text/markdown", "text/plain"}, ".pptx": {"application/vnd.openxmlformats-officedocument.presentationml.presentation", "application/zip"}, ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/zip"}, ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/zip"}}
    if content_type and content_type not in expected.get(ext, set()) | {"application/octet-stream", "binary/octet-stream"} and not content_type.startswith("image/") and ext not in {".xls"}:
        raise ValueError(f"文件类型与扩展名不匹配：{content_type}")
    return name, ext


def _sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _clean_text(value: str) -> str:
    value = re.sub(r"-\s*\n\s*", "", str(value or ""))
    return re.sub(r"(?<!\n)[ \t]*\n[ \t]*", " ", value).strip()


def _block(page: int, index: int, typ: str, text: str, source: str, confidence: float = .9, bbox: list[float] | None = None, **meta: Any) -> dict:
    value = text if typ in {"code", "formula"} else _clean_text(text)
    markdown = value
    if typ == "formula": markdown = f"$$\n{value.strip('$ ').strip()}\n$$" if value.strip() else ""
    if typ == "code": markdown = f"```text\n{value}\n```"
    if typ == "heading": markdown = f"### {value}"
    return {"id": f"p{page}-b{index}", "page_number": page, "type": typ, "raw_text": text, "text": value, "markdown": markdown, "bbox": bbox or [0, 0, 1, 1], "bbox_raw": meta.pop("bbox_raw", bbox or [0, 0, 1, 1]), "confidence": confidence, "source": source, **meta}


def _native_pdf(path: Path) -> list[dict]:
    pages: list[dict] = []
    try:
        import fitz  # type: ignore
        pdf = fitz.open(path)
        for number, page in enumerate(pdf, 1):
            blocks = []
            for index, raw in enumerate(page.get_text("dict").get("blocks", []), 1):
                if raw.get("type") != 0:
                    continue
                text = "\n".join("".join(span.get("text", "") for span in line.get("spans", [])) for line in raw.get("lines", [])).strip()
                if not text:
                    continue
                typ = "heading" if len(text) < 160 and any(span.get("size", 0) >= 14 for line in raw.get("lines", []) for span in line.get("spans", [])) else "text"
                blocks.append(_block(number, index, typ, text, "native_pdf", .96, [max(0, raw.get("bbox", [0, 0, 1, 1])[0] / max(page.rect.width, 1)), max(0, raw.get("bbox", [0, 0, 1, 1])[1] / max(page.rect.height, 1)), min(1, raw.get("bbox", [0, 0, 1, 1])[2] / max(page.rect.width, 1)), min(1, raw.get("bbox", [0, 0, 1, 1])[3] / max(page.rect.height, 1))], bbox_raw=raw.get("bbox")))
            for index, image in enumerate(page.get_image_info(full=True), len(blocks) + 1):
                blocks.append(_block(number, index, "image", "", "native_pdf_image", .9, bbox_raw=image.get("bbox")))
            pages.append({"page": number, "width": page.rect.width, "height": page.rect.height, "blocks": blocks})
        pdf.close()
        return pages
    except Exception:
        try:
            from pypdf import PdfReader  # type: ignore
            reader = PdfReader(str(path))
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                blocks = [_block(number, 1, "text", text, "pypdf", .82)] if text.strip() else []
                pages.append({"page": number, "width": 0, "height": 0, "blocks": blocks})
            return pages
        except Exception as exc:
            raise RuntimeError(f"PDF fallback failed: {exc}") from exc


def _mineru_pdf(path: Path, root: Path) -> list[dict]:
    executable = settings.file_processing_mineru_command
    resolved = shutil.which(executable) or (str(Path(sys.executable).with_name("mineru.exe")) if Path(sys.executable).with_name("mineru.exe").exists() else None)
    if not resolved:
        raise RuntimeError("MinerU command not found")
    out = root / "mineru"
    out.mkdir(parents=True, exist_ok=True)
    cmd = [resolved, "-p", str(path), "-o", str(out), "-m", "auto", "-b", "pipeline", "-f", "true", "-t", "true"]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=3600)
    (out / "mineru.log").write_text((result.stdout or "") + "\n" + (result.stderr or ""), encoding="utf-8")
    if result.returncode != 0:
        raise RuntimeError((result.stderr or result.stdout or "MinerU failed")[-4000:])
    middle = next(iter(out.rglob("*_middle.json")), None)
    if not middle:
        raise RuntimeError("MinerU did not produce middle JSON")
    raw = json.loads(middle.read_text(encoding="utf-8"))
    pages = []
    for number, page in enumerate(raw.get("pdf_info", raw.get("pages", [])), 1):
        blocks = []
        for index, item in enumerate(page.get("para_blocks", page.get("blocks", [])), 1):
            text = item.get("text") or item.get("content") or item.get("latex") or item.get("formula") or ""
            typ = str(item.get("type") or "text").lower()
            if "formula" in typ or "equation" in typ: typ = "formula"
            elif "table" in typ: typ = "table"
            elif "image" in typ or "figure" in typ: typ = "image"
            elif "title" in typ or "heading" in typ: typ = "heading"
            blocks.append(_block(number, index, typ, text, "mineru", float(item.get("confidence", .92)), item.get("bbox"), latex=item.get("latex") or item.get("formula")))
        pages.append({"page": number, "width": page.get("page_width", 0), "height": page.get("page_height", 0), "blocks": blocks})
    if not pages:
        raise RuntimeError("MinerU returned no pages")
    return pages


def _office_pages(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    pages: list[dict] = []
    if ext == ".pptx":
        from pptx import Presentation  # type: ignore
        for number, slide in enumerate(Presentation(str(path)).slides, 1):
            blocks = []
            for index, shape in enumerate(slide.shapes, 1):
                text = getattr(shape, "text", "") or ""
                if text.strip(): blocks.append(_block(number, index, "heading" if getattr(shape, "is_placeholder", False) and getattr(shape, "placeholder_format", None) else "text", text, "python-pptx", .96))
            pages.append({"page": number, "slide_number": number, "width": 0, "height": 0, "blocks": blocks, "title": next((b["text"] for b in blocks if b["type"] == "heading"), "")})
        return pages
    if ext == ".docx":
        from docx import Document  # type: ignore
        blocks = [_block(1, i, "heading" if paragraph.style.name.lower().startswith("heading") else "text", paragraph.text, "python-docx", .95) for i, paragraph in enumerate(Document(str(path)).paragraphs, 1) if paragraph.text.strip()]
        return [{"page": 1, "width": 0, "height": 0, "blocks": blocks}]
    if ext in {".xlsx", ".xls"}:
        from openpyxl import load_workbook  # type: ignore
        book = load_workbook(path, read_only=True, data_only=True)
        for number, sheet in enumerate(book.worksheets, 1):
            lines = [" | ".join(str(value or "") for value in row) for row in sheet.iter_rows(values_only=True)]
            text = "\n".join(line for line in lines if line.strip())
            pages.append({"page": number, "width": 0, "height": 0, "blocks": [_block(number, 1, "table", text, "openpyxl", .95)] if text else [], "sheet": sheet.title})
        return pages
    text = path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".json":
        try: text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except Exception: pass
    return [{"page": 1, "width": 0, "height": 0, "blocks": [_block(1, 1, "text", text, "text", .99)] if text.strip() else []}]


def _office_fallback(path: Path) -> list[dict]:
    """Extract readable XML text when an optional office parser is absent."""
    try:
        with zipfile.ZipFile(path) as archive:
            names = [name for name in archive.namelist() if name.endswith(".xml") and any(token in name for token in ("document", "slide", "sheet", "sharedStrings"))]
            values: list[str] = []
            for name in names:
                try:
                    root = ET.fromstring(archive.read(name))
                    values.extend((node.text or "").strip() for node in root.iter() if node.text and node.text.strip())
                except Exception:
                    continue
            text = "\n".join(dict.fromkeys(values))
            return [{"page": 1, "width": 0, "height": 0, "blocks": [_block(1, 1, "text", text, "zip_xml_fallback", .55)] if text else []}]
    except Exception:
        return [{"page": 1, "width": 0, "height": 0, "blocks": []}]


def _pages_to_chunks(pages: list[dict], filename: str) -> list[dict]:
    chunks: list[dict] = []
    section = ""
    for page in pages:
        for block in page.get("blocks", []):
            content = (block.get("markdown") or block.get("text") or "").strip()
            if not content:
                continue
            if block.get("type") == "heading": section = content[:512]
            chunks.append({"page_number": page.get("page", 0), "section": section, "chunk_type": block.get("type", "text"), "content": content, "metadata": {"filename": filename, "block_id": block.get("id"), "bbox": block.get("bbox"), "bbox_raw": block.get("bbox_raw"), "confidence": block.get("confidence", 0), "source": block.get("source"), "slide_number": page.get("slide_number"), "sheet": page.get("sheet"), "latex": block.get("latex")}})
    return chunks


def _process(db: Session, document: FileProcessingDocument, job: FileProcessingJob) -> None:
    asset = db.get(FileAsset, document.file_asset_id) if document.file_asset_id else db.scalar(select(FileAsset).where(FileAsset.sha256 == document.sha256).order_by(FileAsset.created_at.asc()))
    if not asset:
        raise RuntimeError("原始文件不存在")
    source = settings.upload_path / asset.storage_key
    root = settings.file_processing_path / document.id
    root.mkdir(parents=True, exist_ok=True); (root / "assets").mkdir(exist_ok=True); (root / "logs").mkdir(exist_ok=True)
    target = root / f"source{document.extension}"; shutil.copy2(source, target)
    document.status = "processing"; job.status = "running"; job.stage = "extracting"; job.attempts += 1; job.started_at = now_utc(); db.commit()
    fallback_error = None
    if document.extension == ".pdf":
        try:
            pages = _mineru_pdf(target, root); parser = "mineru"
        except Exception as exc:
            fallback_error = str(exc)[:2000]; pages = _native_pdf(target); parser = "native_pdf_fallback"
    elif document.extension in {".pptx", ".docx", ".xlsx", ".xls", ".txt", ".md", ".markdown", ".csv", ".json"}:
        try:
            pages = _office_pages(target); parser = "native_office"
        except Exception as exc:
            fallback_error = str(exc)[:2000]; pages = _office_fallback(target) if document.extension in {".pptx", ".docx", ".xlsx", ".xls"} else _office_pages(target); parser = "native_office_fallback"
    else:
        pages = [{"page": 1, "width": 0, "height": 0, "blocks": [_block(1, 1, "image", "", "original", .9)]}]; parser = "original_image"
    job.total_pages = len(pages)
    job.processed_pages = 0
    db.commit()
    chunks = _pages_to_chunks(pages, document.filename)
    db.execute(delete(FileProcessingPage).where(FileProcessingPage.document_id == document.id)); db.execute(delete(FileProcessingChunk).where(FileProcessingChunk.document_id == document.id))
    for index, page in enumerate(pages, 1):
        _write_json(root / "pages" / f"page-{page['page']:03d}.json", page)
        db.add(FileProcessingPage(document_id=document.id, page_number=page["page"], width=page.get("width", 0), height=page.get("height", 0), payload=page))
        job.processed_pages = index
        job.progress = round(index / max(len(pages), 1) * 90, 2)
        job.stage = "indexing"
        db.commit()
    for index, chunk in enumerate(chunks):
        db.add(FileProcessingChunk(document_id=document.id, chunk_index=index, page_number=chunk["page_number"], section=chunk["section"], chunk_type=chunk["chunk_type"], content=chunk["content"], metadata_json=chunk["metadata"]))
    markdown = "\n\n".join(f"<!-- page: {page['page']} -->\n\n" + "\n\n".join(block.get("markdown") or block.get("text", "") for block in page.get("blocks", [])) for page in pages)
    (root / "document.md").write_text(markdown, encoding="utf-8")
    (root / "chunks.jsonl").write_text("".join(json.dumps(chunk, ensure_ascii=False) + "\n" for chunk in chunks), encoding="utf-8")
    manifest = {"document_id": document.id, "filename": document.filename, "file_type": document.extension.lstrip("."), "sha256": document.sha256, "pages": len(pages), "successful_pages": len(pages), "failed_pages": 0, "source": parser, "fallback_error": fallback_error, "artifacts": [f"source{document.extension}", "document.md", "manifest.json", "chunks.jsonl"], "page_files": [f"pages/page-{page['page']:03d}.json" for page in pages], "stats": {"chunk_count": len(chunks), "block_count": sum(len(page.get("blocks", [])) for page in pages)}}
    _write_json(root / "manifest.json", manifest)
    document.page_count = len(pages); document.parser = parser; document.artifact_dir = str(root); document.manifest_json = manifest; document.error_message = fallback_error; document.status = "partial_ready" if fallback_error else "ready"
    job.stage = "ready"; job.progress = 100; job.processed_pages = len(pages); job.total_pages = len(pages); job.status = "completed"; job.finished_at = now_utc(); db.commit()
    material_ids = db.scalars(select(CourseMaterial.id).where(CourseMaterial.file_asset_id == document.file_asset_id)).all()
    for material_id in material_ids:
        ingestion = db.scalar(select(CourseMaterialIngestion).where(CourseMaterialIngestion.material_id == material_id))
        if ingestion:
            ingestion.status = document.status; ingestion.parser = parser; ingestion.error_message = fallback_error; ingestion.completed_at = now_utc()
    db.commit()


def process_job(job_id: str) -> None:
    db = SessionLocal()
    retry = False
    try:
        job = db.get(FileProcessingJob, job_id)
        if not job: return
        document = db.get(FileProcessingDocument, job.document_id)
        if not document: return
        try:
            _process(db, document, job)
        except Exception as exc:
            db.rollback(); job = db.get(FileProcessingJob, job_id); document = db.get(FileProcessingDocument, job.document_id) if job else None
            if job and document:
                job.error_message = str(exc)[:4000]
                if job.attempts < job.max_attempts:
                    job.status = "queued"; job.stage = "retry_wait"; document.status = "queued"
                    retry = True
                else:
                    job.status = "failed"; document.status = "failed"; document.error_message = job.error_message; job.finished_at = now_utc()
                material_ids = db.scalars(select(CourseMaterial.id).where(CourseMaterial.file_asset_id == document.file_asset_id)).all() if document else []
                for material_id in material_ids:
                    ingestion = db.scalar(select(CourseMaterialIngestion).where(CourseMaterialIngestion.material_id == material_id))
                    if ingestion: ingestion.status = "failed"; ingestion.error_message = job.error_message
                db.commit()
    finally:
        db.close()
        if retry:
            enqueue_job(job_id)


def enqueue_job(job_id: str) -> None:
    from app.jobs.dispatcher import stage_and_publish

    # The business transaction has already committed before this function is
    # called.  Persisting an outbox row here makes broker outages recoverable
    # without coupling API routes to Celery.
    db = SessionLocal()
    try:
        stage_and_publish(db, "file_processing", (job_id,), f"file-processing:{job_id}")
    finally:
        db.close()


def recover_jobs() -> None:
    db = SessionLocal()
    try:
        jobs = db.scalars(select(FileProcessingJob).where(FileProcessingJob.status.in_(["queued", "running"]))).all()
        job_ids = []
        for job in jobs:
            if job.status == "running": job.status = "queued"
            job_ids.append(job.id)
        db.commit()
        for job_id in job_ids: enqueue_job(job_id)
    finally: db.close()


def ensure_document_for_asset(db: Session, asset: FileAsset, *, target_type: str | None = None, target_id: str | None = None, owner_id: str | None = None, course_id: str | None = None, visibility: str = "owner", metadata: dict | None = None) -> tuple[FileProcessingDocument, FileProcessingJob]:
    # Processing is content-addressed. A second upload of the same bytes can
    # attach a new business binding to the already parsed document.
    document = db.scalar(select(FileProcessingDocument).where(FileProcessingDocument.sha256 == asset.sha256).order_by(FileProcessingDocument.created_at.asc()))
    if document is None:
        document = db.scalar(select(FileProcessingDocument).where(FileProcessingDocument.file_asset_id == asset.id))
    if not document:
        document = FileProcessingDocument(file_asset_id=asset.id, filename=asset.original_name, extension=asset.extension or Path(asset.original_name).suffix.lower(), mime_type=asset.mime_type, sha256=asset.sha256, artifact_dir="", parser="pending", parser_version=settings.file_processing_parser_version)
        db.add(document); db.flush()
    if target_type and target_id:
        binding = db.scalar(select(FileContextBinding).where(FileContextBinding.document_id == document.id, FileContextBinding.target_type == target_type, FileContextBinding.target_id == target_id))
        if not binding: db.add(FileContextBinding(document_id=document.id, target_type=target_type, target_id=target_id, owner_id=owner_id, course_id=course_id, visibility=visibility, metadata_json=metadata or {}))
    job = db.scalar(select(FileProcessingJob).where(FileProcessingJob.document_id == document.id, FileProcessingJob.status.in_(["queued", "running"])).order_by(FileProcessingJob.created_at.desc()))
    if not job and document.status not in {"ready", "partial_ready"}:
        job = FileProcessingJob(document_id=document.id); db.add(job); db.flush()
    if not job:
        job = db.scalar(select(FileProcessingJob).where(FileProcessingJob.document_id == document.id).order_by(FileProcessingJob.created_at.desc()))
    db.commit()
    if job and job.status == "queued": enqueue_job(job.id)
    return document, job


def upload_asset(db: Session, *, uploader_id: str, filename: str, mime_type: str | None, content: bytes, target_type: str | None = None, target_id: str | None = None, owner_id: str | None = None, course_id: str | None = None, visibility: str = "owner", metadata: dict | None = None) -> tuple[FileAsset, FileProcessingDocument, FileProcessingJob]:
    name, ext = validate_upload(filename, mime_type, len(content))
    digest = _sha256(content)
    asset = db.scalar(select(FileAsset).where(FileAsset.uploader_id == uploader_id, FileAsset.sha256 == digest, FileAsset.original_name == name))
    if not asset:
        from app.storage import LocalStorage
        key, size, digest = LocalStorage(settings.upload_path).save(name, content)
        asset = FileAsset(original_name=name, mime_type=mime_type or "application/octet-stream", extension=ext, size_bytes=size, sha256=digest, storage_key=key, uploader_id=uploader_id, course_id=course_id)
        db.add(asset); db.flush()
    document, job = ensure_document_for_asset(db, asset, target_type=target_type, target_id=target_id, owner_id=owner_id or uploader_id, course_id=course_id, visibility=visibility, metadata=metadata)
    return asset, document, job


def bind_document(db: Session, document_id: str, *, target_type: str, target_id: str, owner_id: str | None = None, course_id: str | None = None, visibility: str = "owner", metadata: dict | None = None) -> FileContextBinding:
    binding = db.scalar(select(FileContextBinding).where(FileContextBinding.document_id == document_id, FileContextBinding.target_type == target_type, FileContextBinding.target_id == target_id))
    if binding:
        if metadata:
            binding.metadata_json = {**(binding.metadata_json or {}), **metadata}
            db.commit()
        return binding
    binding = FileContextBinding(document_id=document_id, target_type=target_type, target_id=target_id, owner_id=owner_id, course_id=course_id, visibility=visibility, metadata_json=metadata or {})
    db.add(binding); db.commit(); db.refresh(binding); return binding
